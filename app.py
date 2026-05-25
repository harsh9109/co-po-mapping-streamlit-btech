"""
╔══════════════════════════════════════════════════════════════════════════╗
║   BTech CO-PO Attainment & Academic Analytics System  — v4.0            ║
║   NBA Accreditation Tool | BTech Theory (CCE + End Semester)             ║
╠══════════════════════════════════════════════════════════════════════════╣
║  Theory pattern: 5×CCE (10 marks each) + End Semester (50); 40:60        ║
║  • DI top%, bottom%, threshold — user-configurable                       ║
║  • Unmapped CO averages show "—"                                         ║
║  • PDF A4 portrait · Excel multi-sheet reports · Practical mode intact   ║
╚══════════════════════════════════════════════════════════════════════════╝
"""

import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from io import BytesIO
import warnings
warnings.filterwarnings("ignore")

# reportlab for proper A4 PDF
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm, mm
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                    Spacer, PageBreak, Image, HRFlowable)
    from reportlab.lib.enums import TA_LEFT
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

# ─── Page config ────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="BTech CO-PO Attainment | NBA",
    page_icon="📘",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ══════════════════════════════════════════════════════════════════════════
#  GLOBAL CSS
# ══════════════════════════════════════════════════════════════════════════
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:ital,wght@0,400;0,500;0,600;0,700;0,800;1,400&family=JetBrains+Mono:wght@400;500&display=swap');

:root{
  --ink:#0f172a;--ink-soft:#334155;--muted:#64748b;
  --brand-950:#0b1020;--brand-900:#151d33;--brand-800:#1e293b;
  --primary:#6366f1;--primary-hover:#4f46e5;--primary-glow:rgba(99,102,241,.35);
  --accent:#06b6d4;--accent-soft:#ecfeff;
  --success:#10b981;--warning:#f59e0b;--danger:#ef4444;
  --surface:#ffffff;--surface-2:#f8fafc;--surface-3:#f1f5f9;
  --border:#e2e8f0;--border-strong:#cbd5e1;
  --radius-sm:8px;--radius-md:12px;--radius-lg:16px;--radius-xl:20px;
  --shadow-xs:0 1px 2px rgba(15,23,42,.04);
  --shadow-sm:0 2px 8px rgba(15,23,42,.06);
  --shadow-md:0 8px 24px rgba(15,23,42,.08);
  --shadow-lg:0 16px 48px rgba(15,23,42,.12);
}

*{box-sizing:border-box}
html,body,[class*="css"]{font-family:'Plus Jakarta Sans',system-ui,sans-serif!important;-webkit-font-smoothing:antialiased}
.stApp{
  background:
    radial-gradient(ellipse 80% 50% at 50% -20%,rgba(99,102,241,.09),transparent 55%),
    radial-gradient(ellipse 60% 40% at 100% 0%,rgba(6,182,212,.06),transparent 50%),
    linear-gradient(180deg,#f4f6fb 0%,#eef2f8 100%)!important;
}
.main .block-container{padding:0.4rem 1rem 0.9rem;max-width:100%}

/* Tighten Streamlit default vertical gaps */
.main [data-testid="stVerticalBlock"]{gap:0.4rem!important}
.main [data-testid="stVerticalBlock"]>div{gap:0.4rem!important}
.main .stElementContainer,.main [data-testid="stElementContainer"]{margin-bottom:0.15rem!important}
.main [data-testid="stMarkdownContainer"] p{margin:0 0 0.35rem!important}
.main [data-testid="stMarkdownContainer"] p:empty{display:none!important}
.main [data-testid="column"]{gap:0.5rem!important}
.main hr.sec-div{margin:10px 0!important}
/* Orphan wrapper tags from split HTML blocks */
.main [data-testid="stMarkdownContainer"] div.result-section:empty,
.main [data-testid="stMarkdownContainer"] div.block-card:empty{display:none!important}

/* ── HEADER ── */
.nba-hdr{
  background:linear-gradient(125deg,var(--brand-950) 0%,#1e1b4b 42%,#4338ca 100%);
  border-radius:var(--radius-xl);padding:18px 24px;margin-bottom:14px;
  display:flex;align-items:center;gap:20px;
  box-shadow:var(--shadow-lg),inset 0 1px 0 rgba(255,255,255,.08);
  border:1px solid rgba(255,255,255,.1);position:relative;overflow:hidden;
}
.nba-hdr::before{
  content:'';position:absolute;inset:0;
  background:
    radial-gradient(circle at 85% 15%,rgba(255,255,255,.12) 0%,transparent 42%),
    radial-gradient(circle at 10% 90%,rgba(6,182,212,.15) 0%,transparent 38%);
  pointer-events:none;
}
.nba-hdr::after{
  content:'';position:absolute;right:-80px;bottom:-80px;width:220px;height:220px;
  border-radius:50%;border:1px solid rgba(255,255,255,.06);
  background:transparent;
}
.hdr-logo{
  flex-shrink:0;width:56px;height:56px;border-radius:14px;z-index:1;
  display:flex;align-items:center;justify-content:center;
  background:linear-gradient(145deg,rgba(255,255,255,.18),rgba(255,255,255,.06));
  border:1px solid rgba(255,255,255,.22);
  box-shadow:0 8px 24px rgba(0,0,0,.25);
  font-size:11px;font-weight:800;letter-spacing:.08em;color:#fff;
}
.hdr-text{z-index:1;min-width:0}
.hdr-title{color:#fff;font-size:21px;font-weight:800;margin:0;letter-spacing:-.4px;line-height:1.25}
.hdr-sub{color:rgba(199,210,254,.95);font-size:12.5px;margin:6px 0 0;font-weight:500;line-height:1.45}
.hdr-pills{margin-left:auto;display:flex;gap:7px;flex-wrap:wrap;z-index:1}
.hdr-pill{
  background:rgba(255,255,255,.08);backdrop-filter:blur(8px);
  border:1px solid rgba(255,255,255,.16);color:#f8fafc;
  padding:6px 13px;border-radius:999px;font-size:10.5px;font-weight:700;
  letter-spacing:.35px;white-space:nowrap;
}

.section-head{font-size:18px;font-weight:800;color:var(--ink);letter-spacing:-.35px;margin:0 0 4px}
.section-sub{font-size:12.5px;color:var(--muted);line-height:1.45;margin:0 0 6px}

.ready-strip{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 6px}
.ready-pill{
  display:inline-flex;align-items:center;gap:6px;padding:7px 13px;border-radius:999px;
  font-size:11px;font-weight:700;letter-spacing:.15px;border:1px solid transparent;
  box-shadow:var(--shadow-xs);
}
.ready-pill.ok{background:#ecfdf5;color:#047857;border-color:#a7f3d0}
.ready-pill.warn{background:#fffbeb;color:#b45309;border-color:#fde68a}
.ready-pill.neutral{background:#eef2ff;color:#4338ca;border-color:#c7d2fe}

.empty-state{
  background:var(--surface);border:1.5px dashed var(--border-strong);
  border-radius:var(--radius-xl);padding:16px 18px;margin:4px 0 8px;text-align:center;
  box-shadow:var(--shadow-sm);
}
.empty-title{font-size:18px;font-weight:800;color:var(--ink);letter-spacing:-.25px}
.empty-copy{font-size:13px;color:var(--muted);line-height:1.65;max-width:640px;margin:8px auto 0}

.chart-note,.preview-note{font-size:12.5px;color:var(--muted);line-height:1.65;margin:0 0 10px}

.download-grid{
  display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:12px;margin:8px 0 14px;
}
.download-card{
  background:var(--surface);border:1px solid var(--border);border-radius:var(--radius-lg);
  padding:18px 20px;box-shadow:var(--shadow-sm);transition:transform .2s,box-shadow .2s,border-color .2s;
}
.download-card:hover{
  transform:translateY(-2px);box-shadow:var(--shadow-md);border-color:#c7d2fe;
}
.download-card .dk{font-size:10px;font-weight:700;color:var(--muted);letter-spacing:.65px;text-transform:uppercase}
.download-card .dv{font-size:18px;font-weight:800;color:var(--ink);margin-top:6px;letter-spacing:-.2px}
.download-card .ds{font-size:12.5px;color:var(--muted);line-height:1.6;margin-top:6px}

.stitle{
  font-size:13.5px;font-weight:800;color:var(--ink);
  border-left:3px solid var(--primary);padding-left:12px;
  margin:8px 0 5px;letter-spacing:-.1px;display:flex;align-items:center;gap:8px;flex-wrap:wrap;
}
.stitle:first-child,.block-card>.stitle{margin-top:0!important}
.stitle-meta{font-weight:500;font-size:12px;color:var(--muted);letter-spacing:0}

.mstrip{display:grid;grid-template-columns:repeat(auto-fill,minmax(168px,1fr));gap:10px;margin-bottom:10px}
.mbox{
  background:var(--surface);border-radius:var(--radius-md);padding:12px 14px;
  box-shadow:var(--shadow-sm);border:1px solid var(--border);
  border-top:3px solid var(--ac,var(--primary));transition:transform .2s,box-shadow .2s;
}
.mbox:hover{transform:translateY(-3px);box-shadow:var(--shadow-md)}
.ml{font-size:10px;font-weight:700;color:var(--muted);text-transform:uppercase;letter-spacing:.65px}
.mv{font-size:27px;font-weight:800;color:var(--ink);margin:6px 0 4px;line-height:1;font-variant-numeric:tabular-nums}
.ms{font-size:11.5px;color:var(--muted)}

.awarn{background:#fffbeb;border:1px solid #fde68a;border-left:4px solid var(--warning);
  padding:13px 16px;border-radius:var(--radius-sm);margin:8px 0;font-size:13px;line-height:1.65;color:#78350f}
.agood{background:#ecfdf5;border:1px solid #a7f3d0;border-left:4px solid var(--success);
  padding:13px 16px;border-radius:var(--radius-sm);margin:8px 0;font-size:13px;line-height:1.65;color:#065f46}
.ainfo{background:#eef2ff;border:1px solid #c7d2fe;border-left:4px solid var(--primary);
  padding:13px 16px;border-radius:var(--radius-sm);margin:8px 0;font-size:13px;line-height:1.65;color:#312e81}
.adanger{background:#fff1f2;border:1px solid #fecdd3;border-left:4px solid var(--danger);
  padding:13px 16px;border-radius:var(--radius-sm);margin:8px 0;font-size:13px;line-height:1.65;color:#9f1239}

.nba-wrap{
  overflow-x:auto;border-radius:var(--radius-md);
  box-shadow:var(--shadow-sm);margin:4px 0 8px;border:1px solid var(--border);background:var(--surface);
}
.nba-tbl{width:100%;border-collapse:collapse;font-size:12.5px}
.nba-tbl th{
  background:var(--brand-900);color:#f8fafc;padding:10px 12px;
  text-align:center;font-weight:700;letter-spacing:.2px;
  white-space:nowrap;border-right:1px solid rgba(255,255,255,.08);
}
.nba-tbl th.proc{background:var(--brand-950);text-align:left;min-width:140px}
.nba-tbl td{
  padding:9px 12px;text-align:center;border-bottom:1px solid var(--border);
  border-right:1px solid var(--border);font-variant-numeric:tabular-nums;font-size:12.5px;color:var(--ink-soft);
}
.nba-tbl tr:nth-child(even) td{background:var(--surface-2)}
.nba-tbl tr:hover td{background:#eef2ff;transition:background .15s}
.nba-tbl td.proc{text-align:left;font-weight:700;background:var(--surface-2)!important;
  white-space:nowrap;color:var(--ink)}
.nba-tbl tr.avg-r td{
  background:#e0e7ff!important;font-weight:800;color:var(--ink);
  border-top:2px solid var(--primary);
}
.nba-tbl tr.avg-r td.proc{background:var(--brand-900)!important;color:#fff}
.zero{color:#94a3b8;font-style:italic}
.hi{color:var(--primary-hover);font-weight:800}
.warn{color:var(--danger);font-weight:700}
.good-di{color:var(--success);font-weight:700}
.dash-cell{color:#94a3b8}

.lv0{background:#fee2e2;color:#b91c1c;border:1px solid #fecaca;padding:2px 9px;border-radius:999px;font-size:11px;font-weight:700}
.lv1{background:#fef3c7;color:#92400e;border:1px solid #fde68a;padding:2px 9px;border-radius:999px;font-size:11px;font-weight:700}
.lv2{background:#d1fae5;color:#065f46;border:1px solid #6ee7b7;padding:2px 9px;border-radius:999px;font-size:11px;font-weight:700}
.lv3{background:#e0e7ff;color:#3730a3;border:1px solid #a5b4fc;padding:2px 9px;border-radius:999px;font-size:11px;font-weight:700}

.stTabs [data-baseweb="tab-list"]{
  background:var(--surface);border-radius:var(--radius-md);padding:5px;gap:5px;
  border:1px solid var(--border);box-shadow:var(--shadow-xs);margin-bottom:8px;
}
.stTabs [data-baseweb="tab"]{
  font-size:12.5px;font-weight:700;color:var(--muted);
  padding:10px 18px;border-radius:var(--radius-sm);transition:all .2s;
}
.stTabs [data-baseweb="tab"]:hover{color:var(--ink)}
.stTabs [aria-selected="true"]{
  background:linear-gradient(135deg,var(--brand-900),var(--primary))!important;
  color:#fff!important;box-shadow:0 4px 14px var(--primary-glow)!important;
}

.stButton>button{
  background:linear-gradient(135deg,var(--brand-900),var(--primary))!important;
  color:#fff!important;border:none!important;border-radius:var(--radius-sm)!important;
  font-weight:700!important;padding:10px 26px!important;font-size:13.5px!important;
  box-shadow:0 4px 14px var(--primary-glow)!important;transition:all .22s!important;
}
.stButton>button:hover{
  background:linear-gradient(135deg,#1e1b4b,var(--primary-hover))!important;
  transform:translateY(-2px)!important;box-shadow:0 8px 22px var(--primary-glow)!important;
}
[data-testid="stDownloadButton"]>button{
  background:linear-gradient(135deg,#0f766e,#14b8a6)!important;
  color:#fff!important;border:none!important;border-radius:var(--radius-sm)!important;
  font-weight:700!important;font-size:13.5px!important;
  box-shadow:0 4px 14px rgba(20,184,166,.35)!important;
}
[data-testid="stDownloadButton"]>button:hover{
  background:linear-gradient(135deg,#115e59,#2dd4bf)!important;transform:translateY(-2px)!important;
}

.stNumberInput input,.stTextInput input,.stTextArea textarea{
  border-radius:var(--radius-sm)!important;border:1.5px solid var(--border)!important;
  font-size:13px!important;background:var(--surface)!important;color:var(--ink)!important;
}
.stNumberInput input:focus,.stTextInput input:focus,.stTextArea textarea:focus{
  border-color:var(--primary)!important;box-shadow:0 0 0 3px rgba(99,102,241,.15)!important;
}

div[data-testid="stRadio"]>div{
  background:var(--surface-2);border:1px solid var(--border);border-radius:var(--radius-md);
  padding:6px 8px;gap:6px;
}
div[data-testid="stRadio"] label p{font-weight:600!important;font-size:13px!important}

[data-testid="stDataFrame"],[data-testid="stDataEditor"]{
  border:1px solid var(--border);border-radius:var(--radius-md);overflow:hidden;box-shadow:var(--shadow-xs);
}
[data-testid="stFileUploader"]{
  border:1.5px dashed var(--border-strong);border-radius:var(--radius-md);
  padding:10px;background:rgba(255,255,255,.7);
}
[data-testid="stFileUploader"]:hover{border-color:var(--primary);background:#fafaff}
[data-testid="stVerticalBlockBorderWrapper"]{
  border-radius:var(--radius-lg);border:1px solid var(--border);
  background:var(--surface);box-shadow:var(--shadow-sm);
  padding:0.55rem 0.75rem!important;margin-bottom:0.35rem!important;
}
[data-testid="stVerticalBlockBorderWrapper"] [data-testid="stMarkdownContainer"] p{line-height:1.7;color:var(--ink-soft)}

[data-testid="stPlotlyChart"]{
  border:1px solid var(--border);border-radius:var(--radius-md);
  overflow:hidden;box-shadow:var(--shadow-xs);background:var(--surface);
}

.subj-box{
  background:linear-gradient(125deg,var(--brand-950) 0%,#312e81 55%,#4338ca 100%);
  border-radius:var(--radius-lg);padding:16px 18px;color:#fff;margin-bottom:10px;
  box-shadow:var(--shadow-md);border:1px solid rgba(255,255,255,.1);
}
.subj-box .sc{font-size:10px;font-weight:700;color:#c7d2fe;text-transform:uppercase;letter-spacing:.85px;margin-bottom:5px}
.subj-box .sn{font-size:20px;font-weight:800;margin:0 0 14px;letter-spacing:-.3px}
.subj-box .co-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(280px,1fr));gap:8px;margin-top:8px}
.subj-box .co-item{
  background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.14);
  padding:8px 12px;border-radius:var(--radius-sm);font-size:11.5px;line-height:1.45;
}
.subj-box .co-lbl{font-weight:700;color:#a5b4fc}

.di-config{
  background:linear-gradient(135deg,#fffbeb,#fef3c7);
  border:1px solid #fde68a;border-radius:var(--radius-md);padding:14px 18px;margin:10px 0;
}

.insight-card{
  background:var(--surface);border-radius:var(--radius-md);padding:18px 20px;
  box-shadow:var(--shadow-sm);border:1px solid var(--border);margin:10px 0;
}

.att-bar-wrap{display:flex;align-items:center;gap:10px;margin:4px 0}
.att-bar{height:8px;border-radius:999px;background:var(--surface-3);flex:1;overflow:hidden}
.att-bar-fill{height:100%;border-radius:999px;transition:width .4s}

.sec-div{border:none;border-top:1px solid var(--border);margin:12px 0}

.block-card{
  background:var(--surface);border:1px solid var(--border);border-radius:var(--radius-lg);
  padding:11px 13px;margin:0 0 10px;box-shadow:var(--shadow-xs);
}
.table-block{margin:4px 0 10px}
.tbl-scroll-hint{
  font-size:10px;color:var(--muted);text-align:right;margin-top:4px;letter-spacing:.15px;
}
.nba-wrap::-webkit-scrollbar{height:8px}
.nba-wrap::-webkit-scrollbar-thumb{background:#cbd5e1;border-radius:999px}
.subtable-wrap .nba-wrap{max-width:480px}
.tbl-label{
  font-size:12px;font-weight:700;color:var(--ink-soft);margin:8px 0 4px;letter-spacing:.02em;
}
.insight-card .ic-title{font-size:14px;font-weight:800;color:var(--ink);margin:0 0 8px}
.insight-card .ic-body{font-size:13px;color:var(--ink-soft);line-height:1.65;margin:0}

[data-testid="stWidgetLabel"] p{
  font-weight:600!important;color:var(--ink-soft)!important;font-size:13px!important;
}
.stTabs [data-baseweb="tab-panel"]{padding-top:0;padding-bottom:0}
[data-testid="stPlotlyChart"]{margin-bottom:0.35rem!important}
[data-testid="stPlotlyChart"]>div{min-height:0!important}
.stButton>button[kind="secondary"]{
  background:var(--surface)!important;color:var(--ink-soft)!important;
  border:1.5px solid var(--border)!important;box-shadow:var(--shadow-xs)!important;
}
.stButton>button[kind="secondary"]:hover{
  border-color:var(--primary)!important;color:var(--primary)!important;
  transform:translateY(-1px)!important;box-shadow:var(--shadow-sm)!important;
}
[data-testid="stExpander"]{
  border:1px solid var(--border)!important;border-radius:var(--radius-md)!important;
  background:var(--surface)!important;box-shadow:var(--shadow-xs)!important;
}
div[data-testid="stAlert"]{border-radius:var(--radius-sm)!important}

@media (max-width:860px){
  .nba-hdr{padding:22px 20px;gap:14px;flex-wrap:wrap}
  .hdr-pills{margin-left:0}
}
@media (max-width:560px){
  .hdr-title{font-size:18px}
  .mv{font-size:23px}
}

#MainMenu,footer,header{visibility:hidden}
::-webkit-scrollbar{width:7px;height:7px}
::-webkit-scrollbar-track{background:transparent}
::-webkit-scrollbar-thumb{background:#cbd5e1;border-radius:999px}
::-webkit-scrollbar-thumb:hover{background:#94a3b8}
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════
#  UI HELPERS
# ══════════════════════════════════════════════════════════════════════════

def ui_stitle(label: str) -> None:
    st.markdown(f'<div class="stitle">{label}</div>', unsafe_allow_html=True)


def ui_section(title: str, subtitle: str = "") -> None:
    sub = f'<div class="section-sub">{subtitle}</div>' if subtitle else ""
    st.markdown(f'<div class="section-head">{title}</div>{sub}', unsafe_allow_html=True)


def ui_table_block(inner_html: str) -> str:
    return (
        f'<div class="table-block">{inner_html}'
        f'<div class="tbl-scroll-hint">Scroll horizontally for additional columns</div>'
        f'</div>'
    )


def _attainment_status_met(value: float, threshold: float = 2.0) -> str:
    return "Met" if value >= threshold else "Below target"


def _attainment_status_short(value: float, threshold: float = 2.0) -> str:
    if value <= 0:
        return "—"
    return "On track" if value >= threshold else "Review"


# ══════════════════════════════════════════════════════════════════════════
#  CONSTANTS
# ══════════════════════════════════════════════════════════════════════════
NUM_COS  = 5
NUM_POS  = 11
NUM_PSOS = 3

INTERNAL_EXAMS = ["CCE 1", "CCE 2", "CCE 3", "CCE 4", "CCE 5"]
EXTERNAL_EXAMS = ["End Semester"]
ALL_EXAMS      = INTERNAL_EXAMS + EXTERNAL_EXAMS
PRACTICAL_INTERNAL_EXAM = "Internal"
PRACTICAL_EXTERNAL_EXAM = "External Combined"

DEFAULT_MAX = {
    "CCE 1":        10,
    "CCE 2":        10,
    "CCE 3":        10,
    "CCE 4":        10,
    "CCE 5":        10,
    "End Semester": 50,
}
DEFAULT_INTERNAL_WEIGHT   = 40.0
DEFAULT_UNIVERSITY_WEIGHT = 60.0
DEFAULT_DI_TOP_PCT        = 27.0
DEFAULT_DI_BOT_PCT        = 27.0
DEFAULT_DI_THRESHOLD      = 0.20
DEFAULT_MODE              = "Theory"
DEFAULT_PRACTICAL_INTERNAL_MAX = 25
DEFAULT_PRACTICAL_COMPONENTS = [
    {"name": "Practical", "max_marks": 50},
    {"name": "Term Work", "max_marks": 15},
    {"name": "Oral", "max_marks": 10},
]

DEFAULT_SUBJ_CODE = "304184"
DEFAULT_SUBJ_NAME = "Microcontrollers"
DEFAULT_CO_STMTS  = [
    "Understand the fundamentals of microcontroller and programming",
    "Interface various electronic components with microcontrollers",
    "Analyze the features of PIC 18F XXXX",
    "Describe the programming details in peripheral support",
    "Develop interfacing models according to applications",
]

DEFAULT_COPO = [
    [3, 1, 2, 1, 1, 0, 0, 0, 0, 0, 0],
    [3, 0, 3, 1, 1, 0, 0, 0, 0, 0, 0],
    [3, 3, 1, 1, 1, 0, 1, 0, 0, 0, 0],
    [3, 0, 3, 1, 1, 0, 0, 0, 0, 0, 0],
    [3, 0, 3, 1, 1, 1, 0, 0, 0, 0, 0],
]

DEFAULT_COPSO = [
    [2, 1, 0],
    [2, 1, 0],
    [1, 2, 0],
    [1, 2, 1],
    [1, 1, 2],
]

DEFAULT_COEXAM = [
    [1, 0, 0, 0, 1, 0],
    [1, 0, 0, 0, 1, 0],
    [0, 1, 0, 1, 0, 1],
    [0, 1, 0, 1, 0, 1],
    [0, 0, 1, 0, 0, 1],
]

# Colour constants
NAV = "0A1628"
BLU = "1740AD"
LBL = "EEF3FF"
EVN = "F7F9FD"
DI_LABEL = "DI (H-L)/Max Marks"


# ══════════════════════════════════════════════════════════════════════════
#  CALCULATION FUNCTIONS
# ══════════════════════════════════════════════════════════════════════════

def attainment_level(pct: float) -> int:
    """Levels: 0 &lt;40%; 1 = 40–55%; 2 = 56–70%; 3 &gt;70% (percentage)."""
    if pct < 40:
        return 0
    if pct < 56:
        return 1
    if pct <= 70:
        return 2
    return 3


def exam_attainment(marks: np.ndarray, max_marks: float):
    """NBA formula: att = (0·P + 1·Q + 2·R + 3·S) / N"""
    marks  = np.asarray(marks, dtype=float)
    pcts   = (marks / max_marks) * 100.0
    levels = np.array([attainment_level(p) for p in pcts])
    n      = len(levels)
    if n == 0:
        return 0.0, levels
    P, Q, R, S = [(levels == k).sum() for k in range(4)]
    return float(0*P + 1*Q + 2*R + 3*S) / n, levels


def discrimination_index(marks: np.ndarray, max_marks: float, top_pct: float, bot_pct: float):
    """
    DI = (H_mean − L_mean) / Max Marks
    top_pct, bot_pct in 0–100 (e.g. 27 means top/bottom 27%)
    """
    arr = np.sort(np.asarray(marks, dtype=float))
    n   = len(arr)
    if n == 0:
        return 0.0, 0.0, 0.0, 0, 0
    max_marks = float(max_marks)
    k_top = max(1, int(np.ceil(n * top_pct / 100.0)))
    k_bot = max(1, int(np.ceil(n * bot_pct / 100.0)))
    H  = arr[-k_top:].mean()
    L  = arr[:k_bot].mean()
    DI = (H - L) / max_marks if max_marks > 0 else 0.0
    DI = min(max(float(DI), 0.0), 1.0)
    return round(float(DI), 4), round(float(H), 4), round(float(L), 4), int(k_top), int(k_bot)


def compute_co_attainments(exam_att_dict, coexam_mat, internal_exams, external_exams, all_exams):
    """Compute per-CO internal and external attainments from exam-level data."""
    co_int_att = np.zeros(NUM_COS)
    co_ext_att = np.zeros(NUM_COS)
    for i in range(NUM_COS):
        # Internal
        int_vals = []
        for exam in internal_exams:
            idx = all_exams.index(exam)
            if coexam_mat[i, idx] > 0:
                int_vals.append(exam_att_dict[exam])
        co_int_att[i] = np.mean(int_vals) if int_vals else np.nan

        # External
        ext_vals = []
        for exam in external_exams:
            idx = all_exams.index(exam)
            if coexam_mat[i, idx] > 0:
                ext_vals.append(exam_att_dict[exam])
        co_ext_att[i] = np.mean(ext_vals) if ext_vals else np.nan

    return co_int_att, co_ext_att


def compute_co_finals(co_int_att, co_ext_att, int_ratio, ext_ratio):
    """
    Final CO = int_ratio * Internal + ext_ratio * External
    If a CO has only internal or only external, use what's available (skip NaN side).
    """
    co_finals = np.zeros(NUM_COS)
    for i in range(NUM_COS):
        has_int = not np.isnan(co_int_att[i])
        has_ext = not np.isnan(co_ext_att[i])
        if has_int and has_ext:
            co_finals[i] = int_ratio * co_int_att[i] + ext_ratio * co_ext_att[i]
        elif has_int:
            co_finals[i] = co_int_att[i]
        elif has_ext:
            co_finals[i] = co_ext_att[i]
        else:
            co_finals[i] = 0.0
    return co_finals


def compute_po_pso(co_finals, mapping):
    """
    contrib[i,j] = (co_finals[i] / 3) * mapping[i,j]
    final[j] = mean(contrib[:,j]) over all COs
    """
    num_out = mapping.shape[1]
    contrib = np.zeros((NUM_COS, num_out))
    for i in range(NUM_COS):
        for j in range(num_out):
            contrib[i, j] = (co_finals[i] / 3.0) * mapping[i, j]
    return contrib, contrib.mean(axis=0)


def normalize_weights(iw, uw):
    total = float(iw) + float(uw)
    if total <= 0:
        return 0.40, 0.60
    return float(iw) / total, float(uw) / total


def assessment_labels(mode: str):
    if mode == "Practical":
        return "Internal", "External"
    return "Internal", "University"


def practical_component_col(idx: int) -> str:
    return f"External_Component_{idx+1}"


def practical_component_columns(components):
    return [practical_component_col(i) for i in range(len(components))]


def get_exam_sets(mode: str):
    if mode == "Practical":
        i_exams = [PRACTICAL_INTERNAL_EXAM]
        e_exams = [PRACTICAL_EXTERNAL_EXAM]
    else:
        i_exams = list(INTERNAL_EXAMS)
        e_exams = list(EXTERNAL_EXAMS)
    return i_exams, e_exams, i_exams + e_exams


def build_practical_calc_df(df_raw: pd.DataFrame, component_cols):
    df = df_raw.copy()
    for col in [PRACTICAL_INTERNAL_EXAM] + component_cols:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
    ext_sum = df[component_cols].sum(axis=1) if component_cols else 0
    out = pd.DataFrame({
        "Student Name": df.get("Student Name", pd.Series([f"Std_{i+1}" for i in range(len(df))])),
        PRACTICAL_INTERNAL_EXAM: df[PRACTICAL_INTERNAL_EXAM],
        PRACTICAL_EXTERNAL_EXAM: ext_sum,
    })
    return out


# ══════════════════════════════════════════════════════════════════════════
#  TABLE DATA-FRAME BUILDERS
# ══════════════════════════════════════════════════════════════════════════

def _safe(v, decimals=4):
    """Format float or return '—' for NaN/missing."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    return f"{float(v):.{decimals}f}"


def build_internal_df(co_names, coexam_mat, exam_att_dict, co_int_att, internal_exams, all_exams):
    rows = []
    for exam in internal_exams:
        idx = all_exams.index(exam)
        row = {"Process": exam}
        for i, co in enumerate(co_names):
            row[co] = round(float(exam_att_dict[exam]), 4) if coexam_mat[i, idx] > 0 else "—"
        rows.append(row)
    avg = {"Process": "Average Attainment of CO"}
    for i, co in enumerate(co_names):
        avg[co] = "—" if np.isnan(co_int_att[i]) else round(float(co_int_att[i]), 4)
    rows.append(avg)
    return pd.DataFrame(rows)


def build_external_df(co_names, coexam_mat, exam_att_dict, co_ext_att, external_exams, all_exams):
    rows = []
    for exam in external_exams:
        idx = all_exams.index(exam)
        row = {"Process": exam}
        for i, co in enumerate(co_names):
            row[co] = round(float(exam_att_dict[exam]), 4) if coexam_mat[i, idx] > 0 else "—"
        rows.append(row)
    avg = {"Process": "Average Attainment of CO"}
    for i, co in enumerate(co_names):
        avg[co] = "—" if np.isnan(co_ext_att[i]) else round(float(co_ext_att[i]), 4)
    rows.append(avg)
    return pd.DataFrame(rows)


def build_final_df(co_names, co_int_att, co_ext_att, co_finals, iw, uw, external_label="University"):
    ir, ur = normalize_weights(iw, uw)
    row1 = {"Process": f"{uw:.1f}% of {external_label}"}
    for i, co in enumerate(co_names):
        if np.isnan(co_ext_att[i]):
            row1[co] = "—"
        else:
            row1[co] = round(float(co_ext_att[i]) * ur, 4)

    row2 = {"Process": f"{iw:.1f}% of Internal"}
    for i, co in enumerate(co_names):
        if np.isnan(co_int_att[i]):
            row2[co] = "—"
        else:
            row2[co] = round(float(co_int_att[i]) * ir, 4)

    row3 = {"Process": "CO Attainment"}
    for i, co in enumerate(co_names):
        row3[co] = round(float(co_finals[i]), 4)

    return pd.DataFrame([row1, row2, row3])


def build_di_df(exam_di_dict, top_pct, bot_pct, threshold, internal_exams):
    rows = []
    for exam, d in exam_di_dict.items():
        rows.append({
            "Exam":                               exam,
            "Type":                               "Internal" if exam in internal_exams else "External",
            "N (Students)":                       d["N"],
            f"Top {top_pct:.0f}% Mean (H)":       round(float(d["H"]), 4),
            f"Bottom {bot_pct:.0f}% Mean (L)":    round(float(d["L"]), 4),
            "k (Top)":                            d["k_top"],
            "k (Bot)":                            d["k_bot"],
            DI_LABEL:                             round(float(d["DI"]), 4),
            "Status":                             "Good" if d["DI"] >= threshold else "Low",
        })
    return pd.DataFrame(rows)


def build_po_df(co_names, co_finals, po_contrib, po_finals):
    po_names = [f"PO{j+1}" for j in range(NUM_POS)]
    rows = []
    for i, co in enumerate(co_names):
        row = {"CO": co, "CO Attainment": round(float(co_finals[i]), 4)}
        for j, po in enumerate(po_names):
            row[po] = round(float(po_contrib[i, j]), 4)
        rows.append(row)
    avg = {"CO": "Average", "CO Attainment": ""}
    for j, po in enumerate(po_names):
        avg[po] = round(float(po_finals[j]), 4)
    rows.append(avg)
    return pd.DataFrame(rows)


def build_pso_df(co_names, co_finals, pso_contrib, pso_finals):
    pso_names = [f"PSO{j+1}" for j in range(NUM_PSOS)]
    rows = []
    for i, co in enumerate(co_names):
        row = {"CO": co, "CO Attainment": round(float(co_finals[i]), 4)}
        for j, pso in enumerate(pso_names):
            row[pso] = round(float(pso_contrib[i, j]), 4)
        rows.append(row)
    avg = {"CO": "Average", "CO Attainment": ""}
    for j, pso in enumerate(pso_names):
        avg[pso] = round(float(pso_finals[j]), 4)
    rows.append(avg)
    return pd.DataFrame(rows)


# ══════════════════════════════════════════════════════════════════════════
#  HTML TABLE RENDERERS
# ══════════════════════════════════════════════════════════════════════════

def _v(val, decimals=4):
    """Format cell value for HTML table."""
    if isinstance(val, str):
        if val == "—":
            return '<span class="dash-cell">—</span>'
        return val
    if isinstance(val, float) and np.isnan(val):
        return '<span class="dash-cell">—</span>'
    try:
        f = float(val)
        if f == 0:
            return '<span class="zero">0</span>'
        return f"{f:.{decimals}f}"
    except (TypeError, ValueError):
        return str(val)


def _tbl(headers, rows_data, avg_row_data=None, proc_col=True):
    """Build an HTML NBA-style table."""
    thead_cells = "".join(
        f'<th class="proc">{h}</th>' if (i == 0 and proc_col) else f"<th>{h}</th>"
        for i, h in enumerate(headers)
    )
    tbody = ""
    for row in rows_data:
        cells = "".join(
            f'<td class="proc">{_v(v)}</td>' if (i == 0 and proc_col) else f"<td>{_v(v)}</td>"
            for i, v in enumerate(row)
        )
        tbody += f"<tr>{cells}</tr>"
    if avg_row_data:
        cells = "".join(
            f'<td class="proc">{_v(v)}</td>' if (i == 0 and proc_col) else f"<td>{_v(v)}</td>"
            for i, v in enumerate(avg_row_data)
        )
        tbody += f'<tr class="avg-r">{cells}</tr>'
    return ui_table_block(
        f'<div class="nba-wrap"><table class="nba-tbl"><thead><tr>{thead_cells}</tr></thead>'
        f"<tbody>{tbody}</tbody></table></div>"
    )


def render_internal_table(co_names, coexam_mat, exam_att_dict, co_int_att, internal_exams, all_exams):
    headers = ["Process"] + co_names
    rows = []
    for exam in internal_exams:
        idx = all_exams.index(exam)
        row = [exam]
        for i in range(NUM_COS):
            row.append(round(exam_att_dict[exam], 4) if coexam_mat[i, idx] > 0 else "—")
        rows.append(row)
    avg = ["Avg Attainment of CO"]
    for i in range(NUM_COS):
        avg.append("—" if np.isnan(co_int_att[i]) else round(float(co_int_att[i]), 4))
    return _tbl(headers, rows, avg)


def render_external_table(co_names, coexam_mat, exam_att_dict, co_ext_att, external_exams, all_exams):
    headers = ["Process"] + co_names
    rows = []
    for exam in external_exams:
        idx = all_exams.index(exam)
        row = [exam]
        for i in range(NUM_COS):
            row.append(round(exam_att_dict[exam], 4) if coexam_mat[i, idx] > 0 else "—")
        rows.append(row)
    avg = ["Average Attainment of CO"]
    for i in range(NUM_COS):
        avg.append("—" if np.isnan(co_ext_att[i]) else round(float(co_ext_att[i]), 4))
    return _tbl(headers, rows, avg)


def render_final_table(co_names, co_int_att, co_ext_att, co_finals, iw, uw, external_label="University"):
    ir, ur = normalize_weights(iw, uw)
    headers = ["Process"] + co_names
    rows = []
    r1 = [f"{uw:.1f}% of {external_label}"]
    r2 = [f"{iw:.1f}% of Internal"]
    for i in range(NUM_COS):
        r1.append("—" if np.isnan(co_ext_att[i]) else round(float(co_ext_att[i]) * ur, 4))
        r2.append("—" if np.isnan(co_int_att[i]) else round(float(co_int_att[i]) * ir, 4))
    rows = [r1, r2]
    fin = ["CO Attainment"] + [round(float(v), 4) for v in co_finals]
    return _tbl(headers, rows, fin)


def render_di_table(exam_di_dict, top_pct, bot_pct, threshold, internal_exams):
    headers = ["Exam", "Type", "N", f"Top {top_pct:.0f}% (H)", f"Bot {bot_pct:.0f}% (L)",
               "k(top)", "k(bot)", DI_LABEL, "Status"]
    rows = []
    for exam, d in exam_di_dict.items():
        t    = "Internal" if exam in internal_exams else "External"
        ok   = d["DI"] >= threshold
        stat = '<span class="good-di">Good</span>' if ok else '<span class="warn">Low</span>'
        rows.append([exam, t, d["N"],
                     f'{d["H"]:.4f}', f'{d["L"]:.4f}',
                     d["k_top"], d["k_bot"],
                     f'<span class="{"good-di" if ok else "warn"}">{d["DI"]:.4f}</span>',
                     stat])
    # Don't pass through _v for pre-formatted HTML cells; build directly
    thead = "".join(f'<th class="proc">{h}</th>' if i == 0 else f"<th>{h}</th>"
                    for i, h in enumerate(headers))
    tbody = ""
    for row in rows:
        cells = "".join(f'<td class="proc">{row[0]}</td>' +
                        "".join(f"<td>{v}</td>" for v in row[1:]))
        tbody += f"<tr>{cells}</tr>"
    return ui_table_block(
        f'<div class="nba-wrap"><table class="nba-tbl"><thead><tr>{thead}</tr></thead>'
        f"<tbody>{tbody}</tbody></table></div>"
    )


def render_po_table(co_names, co_finals, po_contrib, po_finals):
    po_names = [f"PO{j+1}" for j in range(NUM_POS)]
    rows = []
    for i, co in enumerate(co_names):
        att = co_finals[i]
        cls = "hi" if att >= 2.0 else "warn"
        att_cell = f'<span class="{cls}">{att:.4f}</span>'
        cells = [co, att_cell] + [
            '<span class="zero">0</span>' if po_contrib[i,j] == 0 else f"{po_contrib[i,j]:.4f}"
            for j in range(NUM_POS)
        ]
        rows.append(cells)
    avg = ["Average", ""] + [
        f'<span class="{"hi" if v >= 2.0 else "warn" if v > 0 else "zero"}">{v:.4f}</span>'
        for v in po_finals
    ]
    # Build raw (already HTML-formatted)
    thead = '<th class="proc">CO</th><th>CO Att.</th>' + "".join(f"<th>{p}</th>" for p in po_names)
    tbody = ""
    for row in rows:
        cells = f'<td class="proc">{row[0]}</td>' + "".join(f"<td>{v}</td>" for v in row[1:])
        tbody += f"<tr>{cells}</tr>"
    avg_cells = '<td class="proc">Average</td>' + "".join(f"<td>{v}</td>" for v in avg[1:])
    tbody += f'<tr class="avg-r">{avg_cells}</tr>'
    return ui_table_block(
        f'<div class="nba-wrap"><table class="nba-tbl"><thead><tr>{thead}</tr></thead>'
        f"<tbody>{tbody}</tbody></table></div>"
    )


def render_pso_table(co_names, co_finals, pso_contrib, pso_finals):
    pso_names = [f"PSO{j+1}" for j in range(NUM_PSOS)]
    rows = []
    for i, co in enumerate(co_names):
        att = co_finals[i]
        cls = "hi" if att >= 2.0 else "warn"
        att_cell = f'<span class="{cls}">{att:.4f}</span>'
        cells = [co, att_cell] + [
            '<span class="zero">0</span>' if pso_contrib[i,j] == 0 else f"{pso_contrib[i,j]:.4f}"
            for j in range(NUM_PSOS)
        ]
        rows.append(cells)
    avg = ["Average", ""] + [
        f'<span class="{"hi" if v >= 2.0 else "warn" if v > 0 else "zero"}">{v:.4f}</span>'
        for v in pso_finals
    ]
    thead = '<th class="proc">CO</th><th>CO Att.</th>' + "".join(f"<th>{p}</th>" for p in pso_names)
    tbody = ""
    for row in rows:
        cells = f'<td class="proc">{row[0]}</td>' + "".join(f"<td>{v}</td>" for v in row[1:])
        tbody += f"<tr>{cells}</tr>"
    avg_cells = '<td class="proc">Average</td>' + "".join(f"<td>{v}</td>" for v in avg[1:])
    tbody += f'<tr class="avg-r">{avg_cells}</tr>'
    return ui_table_block(
        f'<div class="nba-wrap"><table class="nba-tbl"><thead><tr>{thead}</tr></thead>'
        f"<tbody>{tbody}</tbody></table></div>"
    )


# ══════════════════════════════════════════════════════════════════════════
#  EXCEL EXPORT
# ══════════════════════════════════════════════════════════════════════════

def export_excel(R, subj_name, top_pct, bot_pct, di_threshold) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb  = Workbook()
    ws0 = wb.active

    def hdr(ws, r, c, val, fg=NAV, sz=11, bold=True, wrap=True, center=True):
        cell = ws.cell(r, c, val)
        cell.font      = Font(bold=bold, color="FFFFFF", size=sz, name="Calibri")
        cell.fill      = PatternFill("solid", fgColor=fg)
        cell.alignment = Alignment(
            horizontal="center" if center else "left",
            vertical="center", wrap_text=wrap)
        return cell

    def dat(ws, r, c, val, bold=False, even=False, center=True, color=None):
        cell = ws.cell(r, c, val)
        cell.font      = Font(bold=bold, color=color or "000000", size=10, name="Calibri")
        cell.fill      = PatternFill("solid", fgColor="F7F9FD" if even else "FFFFFF")
        cell.alignment = Alignment(
            horizontal="center" if center else "left", vertical="center")
        return cell

    def avg_cell(ws, r, c, val):
        cell = ws.cell(r, c, val)
        cell.font      = Font(bold=True, color="000000", size=10, name="Calibri")
        cell.fill      = PatternFill("solid", fgColor=LBL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        return cell

    thin = Side(style="thin", color="D0D7E8")
    bd   = Border(bottom=thin, right=thin, top=thin, left=thin)

    def apply_bd(ws, r1, r2, c1, c2):
        for row in ws.iter_rows(r1, r2, c1, c2):
            for cell in row:
                cell.border = bd

    def set_col_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def section_title(ws, r, c1, c2, text):
        ws.merge_cells(f"{get_column_letter(c1)}{r}:{get_column_letter(c2)}{r}")
        cell = ws.cell(r, c1, text)
        cell.font      = Font(bold=True, size=12, color="FFFFFF", name="Calibri")
        cell.fill      = PatternFill("solid", fgColor=NAV)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 22
        return r + 1

    def write_df_block(ws, start_r, df, title, avg_label="Average Attainment of CO"):
        """Write a table-block; last row treated as average if it matches avg_label."""
        ncols = len(df.columns)
        r = section_title(ws, start_r, 1, ncols, title)
        # Headers
        for ci, col in enumerate(df.columns, 1):
            c = hdr(ws, r, ci, col, fg=BLU)
            ws.column_dimensions[get_column_letter(ci)].width = max(
                ws.column_dimensions[get_column_letter(ci)].width or 10,
                min(len(str(col)) + 4, 22))
        r += 1
        for di_row, (_, row_data) in enumerate(df.iterrows()):
            even = di_row % 2 == 0
            is_avg = str(row_data.iloc[0]).lower().startswith(("average", "co attainment"))
            for ci, val in enumerate(row_data, 1):
                if is_avg:
                    c = avg_cell(ws, r, ci, val if val != "—" else "")
                    if ci == 1:
                        c.font = Font(bold=True, color="FFFFFF", size=10, name="Calibri")
                        c.fill = PatternFill("solid", fgColor=BLU)
                        c.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    c = dat(ws, r, ci, val if val != "—" else "", even=even,
                            center=(ci != 1), bold=(ci == 1))
            r += 1
        apply_bd(ws, start_r + 1, r - 1, 1, ncols)
        return r + 1

    # ── Sheet 1: Cover ───────────────────────────────────────────────────
    ws0.title = "Cover"
    ws0.sheet_view.showGridLines = False
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 60

    ws0.merge_cells("A1:B1")
    t = ws0.cell(1, 1, "BTech CO-PO Attainment Report  |  NBA Accreditation Format")
    t.font      = Font(bold=True, size=16, color="FFFFFF", name="Calibri")
    t.fill      = PatternFill("solid", fgColor=NAV)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws0.row_dimensions[1].height = 32

    info = [
        ("Subject Code",      R["subj_code"]),
        ("Subject Name",      subj_name),
        ("Assessment Mode",   R.get("mode", "Theory")),
        ("Assessment Pattern", "BTech Theory" if R.get("mode", "Theory") == "Theory" else R.get("mode", "")),
        ("Internal Assessment", "5 CCEs of 10 marks each (CCE 1–CCE 5)" if R.get("mode", "Theory") == "Theory" else "Per practical configuration"),
        ("University Assessment", "End Semester Exam of 50 marks" if R.get("mode", "Theory") == "Theory" else "Combined external components"),
        ("Number of COs",     NUM_COS),
        ("Number of POs",     NUM_POS),
        ("Number of PSOs",    NUM_PSOS),
        ("Total Students",    len(R["df"])),
        ("Internal Weight",   f"{R['iw']:.1f}%"),
        (f"{R.get('external_label', 'University')} Weight", f"{R['uw']:.1f}%"),
        ("DI Top Group",      f"{top_pct:.0f}%"),
        ("DI Bottom Group",   f"{bot_pct:.0f}%"),
        ("DI Threshold",      f"{di_threshold:.2f}"),
    ]
    for ri, (lbl, val) in enumerate(info, 2):
        label_cell = ws0.cell(ri, 1, lbl)
        label_cell.font = Font(bold=True, size=11, name="Calibri")
        label_cell.fill = PatternFill("solid", fgColor=LBL)
        v = ws0.cell(ri, 2, val)
        v.font = Font(size=11, name="Calibri")
        ws0.row_dimensions[ri].height = 18

    r = len(info) + 3
    ws0.cell(r, 1, "CO Statements").font = Font(bold=True, size=12, color="FFFFFF", name="Calibri")
    ws0.cell(r, 1).fill = PatternFill("solid", fgColor=BLU)
    ws0.row_dimensions[r].height = 20
    r += 1
    for i, (co, stmt) in enumerate(zip(R["co_names"], R["co_stmts"])):
        ws0.cell(r, 1, co).font = Font(bold=True, size=10, name="Calibri")
        ws0.cell(r, 1).fill = PatternFill("solid", fgColor="EEF3FF" if i % 2 == 0 else "FFFFFF")
        c = ws0.cell(r, 2, stmt)
        c.font      = Font(size=10, name="Calibri")
        c.alignment = Alignment(wrap_text=True)
        ws0.row_dimensions[r].height = 20
        r += 1

    r += 1
    ws0.cell(r, 1, "Maximum Marks per Exam").font = Font(bold=True, size=12, color="FFFFFF", name="Calibri")
    ws0.cell(r, 1).fill = PatternFill("solid", fgColor=BLU)
    r += 1
    for exam, mark in R["max_marks"].items():
        ws0.cell(r, 1, exam).font = Font(bold=True, size=10, name="Calibri")
        ws0.cell(r, 2, mark).font = Font(size=10, name="Calibri")
        r += 1

    # ── Sheet 2: Student Data ────────────────────────────────────────────
    ws2 = wb.create_sheet("Student Data")
    ws2.sheet_view.showGridLines = False
    cols = ["Student Name"] + R["all_exams"]
    for ci, col in enumerate(cols, 1):
        hdr(ws2, 1, ci, col)
        ws2.column_dimensions[get_column_letter(ci)].width = 16
    ws2.row_dimensions[1].height = 20
    for ri, row in R["df"].iterrows():
        even = ri % 2 == 0
        dat(ws2, ri+2, 1, row.get("Student Name", f"Std_{ri+1}"), even=even, center=False)
        for ci, exam in enumerate(R["all_exams"], 2):
            v = row.get(exam, 0)
            dat(ws2, ri+2, ci, round(float(v), 2) if pd.notna(v) else 0, even=even)
    apply_bd(ws2, 1, len(R["df"])+1, 1, len(cols))

    # ── Sheet 3: CO Attainment ───────────────────────────────────────────
    ws3 = wb.create_sheet("CO Attainment")
    ws3.sheet_view.showGridLines = False
    int_df  = build_internal_df(R["co_names"], R["coexam_matrix"],
                                R["exam_att_dict"], R["co_int_att"],
                                R["internal_exams"], R["all_exams"])
    ext_df  = build_external_df(R["co_names"], R["coexam_matrix"],
                                R["exam_att_dict"], R["co_ext_att"],
                                R["external_exams"], R["all_exams"])
    fin_df  = build_final_df(R["co_names"], R["co_int_att"], R["co_ext_att"],
                             R["co_finals"], R["iw"], R["uw"], R.get("external_label", "University"))
    r = write_df_block(ws3, 1, int_df, "a) Attainment of CO through Internal Assessment")
    r = write_df_block(ws3, r, ext_df, f"b) Attainment of CO through {R.get('external_label', 'University')} Assessment")
    write_df_block(ws3, r, fin_df,
                   f"c) Actual CO Attainment = {R['uw']:.1f}% {R.get('external_label', 'University')} + {R['iw']:.1f}% Internal")

    # ── Sheet 4: Discrimination Index ────────────────────────────────────
    ws4 = wb.create_sheet("Discrimination Index")
    ws4.sheet_view.showGridLines = False
    di_df = build_di_df(R["exam_di_dict"], top_pct, bot_pct, di_threshold, R["internal_exams"])
    section_title(ws4, 1, 1, len(di_df.columns),
                  f"Discrimination Index  |  Top {top_pct:.0f}% / Bottom {bot_pct:.0f}%  |  Threshold = {di_threshold:.2f}")
    for ci, col in enumerate(di_df.columns, 1):
        hdr(ws4, 2, ci, col, fg=BLU)
        ws4.column_dimensions[get_column_letter(ci)].width = max(len(str(col)) + 3, 14)
    ws4.row_dimensions[2].height = 20
    for ri, (_, row) in enumerate(di_df.iterrows(), 3):
        even = ri % 2 == 0
        for ci, val in enumerate(row, 1):
            clean_val = str(val)
            c = dat(ws4, ri, ci, clean_val, even=even, center=(ci != 1))
            is_di = (ci == len(di_df.columns) - 1)
            if is_di:
                di_val = float(di_df.iloc[ri-3][DI_LABEL])
                c.font = Font(bold=True, size=10, name="Calibri",
                              color="16A34A" if di_val >= di_threshold else "DC2626")
    apply_bd(ws4, 2, 2+len(di_df), 1, len(di_df.columns))

    # ── Sheet 5: PO Attainment ───────────────────────────────────────────
    ws5 = wb.create_sheet("PO Attainment")
    ws5.sheet_view.showGridLines = False
    po_df = build_po_df(R["co_names"], R["co_finals"],
                        R["po_contrib"], R["po_finals"])
    write_df_block(ws5, 1, po_df, "PO Attainment  [Formula: (CO_att/3) × mapping weight  |  Average over all COs]")

    # ── Sheet 6: PSO Attainment ──────────────────────────────────────────
    ws6 = wb.create_sheet("PSO Attainment")
    ws6.sheet_view.showGridLines = False
    pso_df = build_pso_df(R["co_names"], R["co_finals"],
                          R["pso_contrib"], R["pso_finals"])
    write_df_block(ws6, 1, pso_df, "PSO Attainment  [Formula: (CO_att/3) × mapping weight  |  Average over all COs]")

    # ── Sheet 7: Summary ────────────────────────────────────────────────
    ws7 = wb.create_sheet("Summary")
    ws7.sheet_view.showGridLines = False
    ws7.column_dimensions["A"].width = 20
    for ci in range(2, 7):
        ws7.column_dimensions[get_column_letter(ci)].width = 14

    section_title(ws7, 1, 1, 5, "CO Attainment Summary")
    for ci, h in enumerate(["CO", "Internal", R.get("external_label", "University"), "Final", "Target Met"], 1):
        hdr(ws7, 2, ci, h, fg=BLU)
    for ri, co in enumerate(R["co_names"], 3):
        i = ri - 3
        even = i % 2 == 0
        int_v = R["co_int_att"][i]
        ext_v = R["co_ext_att"][i]
        fin_v = R["co_finals"][i]
        dat(ws7, ri, 1, co, bold=True, even=even, center=False)
        dat(ws7, ri, 2, "—" if np.isnan(int_v) else round(float(int_v), 4), even=even)
        dat(ws7, ri, 3, "—" if np.isnan(ext_v) else round(float(ext_v), 4), even=even)
        dat(ws7, ri, 4, round(float(fin_v), 4), even=even,
            color="1740AD" if fin_v >= 2.0 else "DC2626", bold=True)
        dat(ws7, ri, 5, "Yes" if fin_v >= 2.0 else "No", even=even,
            color="16A34A" if fin_v >= 2.0 else "DC2626", bold=True)
    apply_bd(ws7, 2, 2+NUM_COS, 1, 5)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════
#  PDF EXPORT (A4 PORTRAIT — reportlab)
# ══════════════════════════════════════════════════════════════════════════

def _rl_color(hex_str):
    h = hex_str.lstrip("#")
    return colors.Color(int(h[0:2],16)/255, int(h[2:4],16)/255, int(h[4:6],16)/255)

C_NAV  = _rl_color("061231")
C_BLU  = _rl_color("1740AD")
C_LBL  = _rl_color("EEF3FF")
C_EVN  = _rl_color("F7F9FD")
C_RED  = _rl_color("DC2626")
C_GRN  = _rl_color("16A34A")
C_WHT  = colors.white
C_BLK  = colors.black
C_GREY = _rl_color("D0D7E8")


def _rl_table(data, col_widths, highlight_last=True):
    """Build a formatted reportlab Table from list-of-lists data."""
    style = [
        ("BACKGROUND",   (0, 0), (-1, 0),    C_NAV),
        ("TEXTCOLOR",    (0, 0), (-1, 0),    C_WHT),
        ("FONTNAME",     (0, 0), (-1, 0),    "Helvetica-Bold"),
        ("FONTSIZE",     (0, 0), (-1, 0),    8),
        ("ALIGN",        (0, 0), (-1, -1),   "CENTER"),
        ("ALIGN",        (0, 1), (0, -1),    "LEFT"),
        ("VALIGN",       (0, 0), (-1, -1),   "MIDDLE"),
        ("ROWBACKGROUND",(0, 1), (-1, -1),   [C_WHT, C_EVN]),
        ("GRID",         (0, 0), (-1, -1),   0.4, C_GREY),
        ("FONTNAME",     (0, 1), (-1, -1),   "Helvetica"),
        ("FONTSIZE",     (0, 1), (-1, -1),   7.5),
        ("LEFTPADDING",  (0, 0), (-1, -1),   5),
        ("RIGHTPADDING", (0, 0), (-1, -1),   5),
        ("TOPPADDING",   (0, 0), (-1, -1),   4),
        ("BOTTOMPADDING",(0, 0), (-1, -1),   4),
        ("BACKGROUND",   (0, 0), (0, -1),    C_BLU),
        ("TEXTCOLOR",    (0, 1), (0, -1),    C_WHT),
        ("FONTNAME",     (0, 1), (0, -1),    "Helvetica-Bold"),
    ]
    if highlight_last and len(data) > 2:
        style += [
            ("BACKGROUND",  (0, -1), (-1, -1), C_LBL),
            ("FONTNAME",    (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE",    (0, -1), (-1, -1), 7.5),
            ("BACKGROUND",  (0, -1), (0, -1),  C_BLU),
            ("TEXTCOLOR",   (0, -1), (0, -1),  C_WHT),
            ("LINEABOVE",   (0, -1), (-1, -1), 1.5, C_BLU),
        ]
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle(style))
    return tbl


def _section_para(text, styles):
    style = ParagraphStyle("sec", parent=styles["Normal"],
                           fontSize=11, fontName="Helvetica-Bold",
                           textColor=C_NAV, spaceAfter=4, spaceBefore=10,
                           borderPadding=(4, 0, 4, 10),
                           borderColor=C_BLU, borderWidth=0)
    return Paragraph(f"<b>{text}</b>", style)


def export_pdf(R, subj_name, top_pct, bot_pct, di_threshold) -> BytesIO:
    if not REPORTLAB_OK:
        return _fallback_pdf(R, subj_name)

    buf = BytesIO()
    styles = getSampleStyleSheet()
    W, H = A4   # 595 × 842 points
    margin = 1.8 * cm
    usable_w = W - 2 * margin

    title_style = ParagraphStyle("title", parent=styles["Normal"],
        fontSize=18, fontName="Helvetica-Bold", textColor=C_NAV,
        spaceAfter=6, alignment=TA_LEFT)
    sub_style = ParagraphStyle("sub", parent=styles["Normal"],
        fontSize=10, fontName="Helvetica", textColor=_rl_color("334155"),
        spaceAfter=14, alignment=TA_LEFT)
    body_style = ParagraphStyle("body", parent=styles["Normal"],
        fontSize=9, fontName="Helvetica", textColor=C_BLK,
        spaceAfter=4, leading=14)
    bold_style = ParagraphStyle("boldb", parent=styles["Normal"],
        fontSize=9.5, fontName="Helvetica-Bold", textColor=C_NAV,
        spaceAfter=3, spaceBefore=8)

    def page_header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(C_NAV)
        canvas.rect(margin, H - margin + 2*mm, usable_w, 0.3*mm, fill=1, stroke=0)
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(C_NAV)
        canvas.drawString(margin, H - margin + 4*mm,
            f"{R['subj_code']} | {subj_name}  —  BTech CO-PO Attainment (NBA)")
        canvas.drawRightString(W - margin, H - margin + 4*mm, "v4 BTech")
        # Footer
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(_rl_color("94a3b8"))
        canvas.drawString(margin, 14*mm,
            "CO-PO Attainment · BTech · NBA Tool")
        canvas.drawRightString(W - margin, 14*mm, f"Page {doc.page}")
        canvas.setFillColor(C_BLU)
        canvas.rect(margin, 12*mm, usable_w, 0.3*mm, fill=1, stroke=0)
        canvas.restoreState()

    story = []

    # ─── COVER PAGE ──────────────────────────────────────────────────────
    story.append(Spacer(1, 1.5*cm))
    story.append(Paragraph("BTech CO-PO Attainment Report", title_style))
    story.append(Paragraph(
        f"{R['subj_code']}  |  {subj_name}  |  COs: {NUM_COS}  |  POs: {NUM_POS}  |  PSOs: {NUM_PSOS}",
        sub_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=C_BLU))
    story.append(Spacer(1, 0.35*cm))
    if R.get("mode", "Theory") == "Theory":
        story.append(Paragraph(
            "<b>Assessment Pattern:</b> BTech Theory<br/>"
            "<b>Internal Assessment:</b> 5 CCEs of 10 marks each<br/>"
            "<b>University Assessment:</b> End Semester Exam of 50 marks<br/>"
            "<b>Internal Weight:</b> 40% &nbsp;|&nbsp; <b>University Weight:</b> 60%",
            body_style))
        story.append(Spacer(1, 0.45*cm))

    meta = [
        ["Total Students", str(len(R["df"])), "Internal Weight", f"{R['iw']:.1f}%"],
        [f"{R.get('external_label', 'University')} Weight", f"{R['uw']:.1f}%", "DI Top Group", f"{top_pct:.0f}%"],
        ["DI Bottom Group", f"{bot_pct:.0f}%", "DI Threshold", f"{di_threshold:.2f}"],
    ]
    meta_tbl = Table(meta, colWidths=[4*cm, 4*cm, 4*cm, 4*cm])
    meta_tbl.setStyle(TableStyle([
        ("GRID",        (0,0),(-1,-1), 0.5, C_GREY),
        ("FONTNAME",    (0,0),(0,-1),  "Helvetica-Bold"),
        ("FONTNAME",    (2,0),(2,-1),  "Helvetica-Bold"),
        ("FONTSIZE",    (0,0),(-1,-1), 9),
        ("ROWBACKGROUND",(0,0),(-1,-1),[C_EVN, C_WHT]),
        ("ALIGN",       (0,0),(-1,-1),"LEFT"),
        ("TOPPADDING",  (0,0),(-1,-1), 5),
        ("BOTTOMPADDING",(0,0),(-1,-1),5),
        ("LEFTPADDING", (0,0),(-1,-1), 8),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 0.8*cm))

    story.append(Paragraph("Course Outcomes", bold_style))
    co_data = [["CO", "Statement"]]
    for co, stmt in zip(R["co_names"], R["co_stmts"]):
        co_data.append([co, stmt])
    co_tbl = Table(co_data, colWidths=[3.5*cm, usable_w - 3.5*cm])
    co_tbl.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(-1,0),  C_BLU),
        ("TEXTCOLOR",    (0,0),(-1,0),  C_WHT),
        ("FONTNAME",     (0,0),(-1,0),  "Helvetica-Bold"),
        ("FONTNAME",     (0,1),(-1,-1), "Helvetica"),
        ("FONTNAME",     (0,1),(0,-1),  "Helvetica-Bold"),
        ("FONTSIZE",     (0,0),(-1,-1), 9),
        ("ROWBACKGROUND",(0,1),(-1,-1), [C_WHT, C_EVN]),
        ("GRID",         (0,0),(-1,-1), 0.4, C_GREY),
        ("ALIGN",        (0,0),(-1,-1), "LEFT"),
        ("VALIGN",       (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",   (0,0),(-1,-1), 5),
        ("BOTTOMPADDING",(0,0),(-1,-1), 5),
        ("LEFTPADDING",  (0,0),(-1,-1), 6),
    ]))
    story.append(co_tbl)
    story.append(Spacer(1, 0.45*cm))
    story.append(_section_para("Student Data Summary", styles))
    if R.get("mode", "Theory") == "Theory":
        sum_rows = [
            ["Description", "Value"],
            ["Students analysed", str(len(R["df"]))],
            ["Marks columns", "Student Name; " + "; ".join(R["all_exams"])],
            ["Maximum marks", "CCE 1–CCE 5: 10 each; End Semester: 50 (configurable in app)"],
        ]
    else:
        sum_rows = [
            ["Description", "Value"],
            ["Students analysed", str(len(R["df"]))],
            ["Assessment mode", "Practical — internal + combined external components"],
        ]
    story.append(_rl_table(sum_rows, [4.5*cm, usable_w - 4.5*cm], highlight_last=False))
    story.append(PageBreak())

    # ─── ATTAINMENT LEVEL REFERENCE ──────────────────────────────────────
    story.append(_section_para("Attainment Level Reference", styles))
    lvl_data = [["Level", "Percentage", "Interpretation"],
                ["0", "< 40%", "Not Achieved"],
                ["1", "40% – 55%", "Partially Achieved"],
                ["2", "56% – 70%", "Achieved"],
                ["3", "> 70%", "Highly Achieved"]]
    lvl_tbl = Table(lvl_data, colWidths=[3*cm, 4*cm, 6*cm])
    lvl_tbl.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(-1,0),  C_NAV),
        ("TEXTCOLOR",    (0,0),(-1,0),  C_WHT),
        ("FONTNAME",     (0,0),(-1,0),  "Helvetica-Bold"),
        ("FONTNAME",     (0,1),(-1,-1), "Helvetica"),
        ("FONTSIZE",     (0,0),(-1,-1), 9),
        ("GRID",         (0,0),(-1,-1), 0.4, C_GREY),
        ("ALIGN",        (0,0),(-1,-1), "CENTER"),
        ("TOPPADDING",   (0,0),(-1,-1), 5),
        ("BOTTOMPADDING",(0,0),(-1,-1), 5),
    ]))
    story.append(lvl_tbl)
    story.append(Spacer(1, 0.5*cm))

    # ─── HELPER: df → reportlab table ────────────────────────────────────
    def df_to_rl(df, title, col_widths=None, highlight_last=True):
        story.append(_section_para(title, styles))
        data = [list(df.columns)] + df.fillna("—").astype(str).values.tolist()
        if col_widths is None:
            col_widths = [usable_w / len(df.columns)] * len(df.columns)
            col_widths[0] = min(col_widths[0] * 1.6, 4.5*cm)
        story.append(_rl_table(data, col_widths, highlight_last))
        story.append(Spacer(1, 0.4*cm))

    # ─── CO ATTAINMENT (INTERNAL) ────────────────────────────────────────
    int_df = build_internal_df(R["co_names"], R["coexam_matrix"],
                               R["exam_att_dict"], R["co_int_att"],
                               R["internal_exams"], R["all_exams"])
    cw_int = [3.8*cm] + [(usable_w - 3.8*cm) / NUM_COS] * NUM_COS
    df_to_rl(int_df, "a) Attainment of CO through Internal Assessment", cw_int)

    ext_df = build_external_df(R["co_names"], R["coexam_matrix"],
                               R["exam_att_dict"], R["co_ext_att"],
                               R["external_exams"], R["all_exams"])
    df_to_rl(ext_df, f"b) Attainment of CO through {R.get('external_label', 'University')} Assessment", cw_int)

    fin_df = build_final_df(R["co_names"], R["co_int_att"], R["co_ext_att"],
                            R["co_finals"], R["iw"], R["uw"], R.get("external_label", "University"))
    df_to_rl(fin_df,
             f"c) CO Attainment = {R['uw']:.1f}% {R.get('external_label', 'University')} + {R['iw']:.1f}% Internal",
             cw_int)
    story.append(PageBreak())

    # ─── DISCRIMINATION INDEX ────────────────────────────────────────────
    di_df = build_di_df(R["exam_di_dict"], top_pct, bot_pct, di_threshold, R["internal_exams"])
    di_cw = [3.2*cm, 2.2*cm, 1.5*cm, 2.5*cm, 2.5*cm, 1.5*cm, 1.5*cm, 2.5*cm, 2*cm]
    story.append(_section_para(
        f"Discrimination Index  |  Top {top_pct:.0f}% / Bot {bot_pct:.0f}%  |  Threshold: {di_threshold:.2f}",
        styles))
    di_data = [list(di_df.columns)] + di_df.values.tolist()
    di_tbl = Table(di_data, colWidths=di_cw, repeatRows=1)
    di_style = [
        ("BACKGROUND",  (0,0),(-1,0),  C_NAV),
        ("TEXTCOLOR",   (0,0),(-1,0),  C_WHT),
        ("FONTNAME",    (0,0),(-1,0),  "Helvetica-Bold"),
        ("FONTSIZE",    (0,0),(-1,-1), 7.5),
        ("FONTNAME",    (0,1),(-1,-1), "Helvetica"),
        ("ALIGN",       (0,0),(-1,-1), "CENTER"),
        ("ALIGN",       (0,1),(0,-1),  "LEFT"),
        ("GRID",        (0,0),(-1,-1), 0.4, C_GREY),
        ("TOPPADDING",  (0,0),(-1,-1), 4),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING", (0,0),(-1,-1), 4),
    ]
    for ri, (exam, d) in enumerate(R["exam_di_dict"].items(), 1):
        ok = d["DI"] >= di_threshold
        fc = C_GRN if ok else C_RED
        di_style.append(("TEXTCOLOR", (-2, ri), (-1, ri), fc))
        di_style.append(("FONTNAME",  (-2, ri), (-1, ri), "Helvetica-Bold"))
        if ri % 2 == 0:
            di_style.append(("BACKGROUND", (0, ri), (-1, ri), C_EVN))
    di_tbl.setStyle(TableStyle(di_style))
    story.append(di_tbl)
    story.append(Spacer(1, 0.5*cm))
    story.append(PageBreak())

    # ─── PO ATTAINMENT ───────────────────────────────────────────────────
    po_df = build_po_df(R["co_names"], R["co_finals"],
                        R["po_contrib"], R["po_finals"])
    po_cw = [2.8*cm, 1.8*cm] + [(usable_w - 4.6*cm) / NUM_POS] * NUM_POS
    df_to_rl(po_df, "PO Attainment  [contrib = (CO_att/3) × mapping,  avg over all COs]", po_cw)

    pso_df = build_pso_df(R["co_names"], R["co_finals"],
                          R["pso_contrib"], R["pso_finals"])
    pso_cw = [3*cm, 2*cm] + [(usable_w - 5*cm) / NUM_PSOS] * NUM_PSOS
    df_to_rl(pso_df, "PSO Attainment  [same formula]", pso_cw)
    story.append(PageBreak())

    # --- KEY SUMMARY --------------------------------------------------------
    co_ok = int((R["co_finals"] >= 2.0).sum())
    avg_co = float(np.mean(R["co_finals"])) if len(R["co_finals"]) else 0.0
    po_v = R["po_finals"]
    avg_po = float(po_v[po_v > 0].mean()) if (po_v > 0).any() else 0.0
    low_di = sum(1 for d in R["exam_di_dict"].values() if d["DI"] < di_threshold)
    story.append(_section_para("Key Summary", styles))
    summary_data = [
        ["Metric", "Value"],
        ["Average CO Attainment", f"{avg_co:.3f} / 3.000"],
        ["COs at Target", f"{co_ok} / {NUM_COS} (target >= 2.0)"],
        ["Average PO Attainment", f"{avg_po:.3f} / 3.000"],
        ["Low DI Exams", f"{low_di} / {len(R['all_exams'])} below {di_threshold:.2f}"],
    ]
    story.append(_rl_table(summary_data, [7*cm, usable_w - 7*cm], highlight_last=False))
    story.append(PageBreak())

    # --- CHARTS: one clear A4 section per visual ---------------------------
    def fig_to_img(fig):
        buf2 = BytesIO()
        fig.savefig(buf2, format="png", dpi=170, bbox_inches="tight",
                    facecolor="white")
        plt.close(fig)
        buf2.seek(0)
        return buf2

    def append_chart(title, fig, max_height=17.5*cm):
        width = usable_w
        height = min(width * (fig.get_figheight() / fig.get_figwidth()), max_height)
        story.append(_section_para(title, styles))
        story.append(Image(fig_to_img(fig), width=width, height=height))
        story.append(PageBreak())

    colors_co = [("#1740AD" if v >= 2.5 else "#22c55e" if v >= 2 else
                  "#f59e0b" if v >= 1.5 else "#ef4444") for v in R["co_finals"]]

    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    ax.bar(R["co_names"], np.round(R["co_finals"], 3), color=colors_co)
    ax.axhline(2.0, linestyle="--", color="#ef4444", lw=1.3, label="Target 2.0")
    ax.axhline(avg_co, linestyle=":", color="#1740AD", lw=1.5, label=f"CO Avg {avg_co:.3f}")
    ax.set_title("Final CO Attainment", fontsize=13, fontweight="bold")
    ax.set_ylabel("Attainment")
    ax.set_ylim(0, 3.4)
    ax.grid(axis="y", ls=":", alpha=0.35)
    ax.legend(frameon=False, fontsize=9)
    for i, v in enumerate(R["co_finals"]):
        ax.text(i, v + 0.05, f"{v:.3f}", ha="center", fontsize=9)
    plt.tight_layout(pad=1.8)
    append_chart("Chart 1: Final CO Attainment", fig)

    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    x = np.arange(NUM_COS)
    w = 0.36
    int_plot = np.where(np.isnan(R["co_int_att"]), 0, R["co_int_att"])
    ext_plot = np.where(np.isnan(R["co_ext_att"]), 0, R["co_ext_att"])
    ax.bar(x - w/2, np.round(int_plot, 3), width=w, color="#5c6bc0",
           label=f"Internal ({R['iw']:.0f}%)")
    ax.bar(x + w/2, np.round(ext_plot, 3), width=w, color="#0891b2",
           label=f"{R.get('external_label', 'University')} ({R['uw']:.0f}%)")
    ax.set_xticks(x)
    ax.set_xticklabels(R["co_names"])
    ax.set_title(f"Internal vs {R.get('external_label', 'University')} Attainment", fontsize=13, fontweight="bold")
    ax.set_ylabel("Attainment")
    ax.set_ylim(0, 3.4)
    ax.grid(axis="y", ls=":", alpha=0.35)
    ax.legend(frameon=False, fontsize=9)
    plt.tight_layout(pad=1.8)
    append_chart(f"Chart 2: Internal vs {R.get('external_label', 'University')} Attainment", fig)

    po_names = [f"PO{j+1}" for j in range(NUM_POS)]
    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    ax.bar(po_names, np.round(R["po_finals"], 3),
           color=[("#1740AD" if v >= 2.5 else "#22c55e" if v >= 2 else
                   "#f59e0b" if v >= 1 else "#ef4444") for v in R["po_finals"]])
    ax.axhline(2.0, ls="--", color="#ef4444", lw=1.3, label="Target 2.0")
    ax.set_title("PO Attainment", fontsize=13, fontweight="bold")
    ax.set_ylabel("Attainment")
    ax.set_ylim(0, 3.4)
    ax.tick_params(axis="x", rotation=30)
    ax.grid(axis="y", ls=":", alpha=0.35)
    ax.legend(frameon=False, fontsize=9)
    for i, v in enumerate(R["po_finals"]):
        ax.text(i, v + 0.05, f"{v:.2f}", ha="center", fontsize=8)
    plt.tight_layout(pad=1.8)
    append_chart("Chart 3: PO Attainment", fig)

    pso_names = [f"PSO{j+1}" for j in range(NUM_PSOS)]
    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    ax.bar(pso_names, np.round(R["pso_finals"], 3),
           color=["#1740AD", "#0f766e", "#b45309"][:NUM_PSOS])
    ax.axhline(2.0, ls="--", color="#ef4444", lw=1.3, label="Target 2.0")
    ax.set_title("PSO Attainment", fontsize=13, fontweight="bold")
    ax.set_ylabel("Attainment")
    ax.set_ylim(0, 3.4)
    ax.grid(axis="y", ls=":", alpha=0.35)
    ax.legend(frameon=False, fontsize=9)
    for i, v in enumerate(R["pso_finals"]):
        ax.text(i, v + 0.05, f"{v:.3f}", ha="center", fontsize=10)
    plt.tight_layout(pad=1.8)
    append_chart("Chart 4: PSO Attainment", fig)

    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    di_vals = [R["exam_di_dict"][e]["DI"] for e in R["all_exams"]]
    di_clrs = ["#22c55e" if d >= di_threshold else "#ef4444" for d in di_vals]
    ax.bar(R["all_exams"], di_vals, color=di_clrs)
    ax.axhline(di_threshold, ls="--", color="#f59e0b", lw=1.5,
               label=f"Threshold ({di_threshold:.2f})")
    ax.set_title("Discrimination Index per Exam", fontsize=13, fontweight="bold")
    ax.set_ylabel("DI")
    ax.set_ylim(0, max(max(di_vals)*1.35, di_threshold*2))
    ax.grid(axis="y", ls=":", alpha=0.35)
    ax.legend(frameon=False, fontsize=9)
    for i, v in enumerate(di_vals):
        ax.text(i, v + 0.002, f"{v:.4f}", ha="center", fontsize=8)
    plt.tight_layout(pad=1.8)
    append_chart("Chart 5: Discrimination Index", fig)

    fig, ax = plt.subplots(figsize=(7.4, 4.8))
    fig.patch.set_facecolor("white")
    heat = np.round(R["po_contrib"], 3)
    im = ax.imshow(heat, cmap="Blues", aspect="auto")
    ax.set_xticks(np.arange(NUM_POS))
    ax.set_xticklabels(po_names, rotation=45, ha="right")
    ax.set_yticks(np.arange(NUM_COS))
    ax.set_yticklabels(R["co_names"])
    ax.set_title("CO to PO Contribution Heatmap", fontsize=13, fontweight="bold")
    for i in range(NUM_COS):
        for j in range(NUM_POS):
            ax.text(j, i, f"{heat[i, j]:.2f}", ha="center", va="center", fontsize=7,
                    color="white" if heat[i, j] > heat.max() * 0.55 else "#061231")
    fig.colorbar(im, ax=ax, fraction=0.035, pad=0.02, label="Contribution")
    plt.tight_layout(pad=1.8)
    append_chart("Chart 6: CO to PO Contribution Heatmap", fig)

    # --- BUILD PDF ---------------------------------------------------------
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=margin, rightMargin=margin,
        topMargin=margin + 0.8*cm, bottomMargin=margin + 0.6*cm,
    )
    doc.build(story, onFirstPage=page_header_footer, onLaterPages=page_header_footer)
    buf.seek(0)
    return buf


def _fallback_pdf(R, subj_name) -> BytesIO:
    """Simple matplotlib PDF if reportlab not available."""
    from matplotlib.backends.backend_pdf import PdfPages
    buf = BytesIO()
    with PdfPages(buf) as pdf:
        fig, ax = plt.subplots(figsize=(8.27, 11.69))
        ax.axis("off")
        ax.text(.5, .95, "CO-PO Attainment Report", ha="center", fontsize=18,
                fontweight="bold", transform=ax.transAxes)
        ax.text(.5, .88, f"{R['subj_code']} | {subj_name}", ha="center",
                fontsize=12, transform=ax.transAxes, color="#334155")
        ax.text(.1, .82, "Install reportlab for a full A4 PDF:\n  pip install reportlab",
                fontsize=11, transform=ax.transAxes, color="#7c3aed")
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════
#  PLOTLY CHARTS
# ══════════════════════════════════════════════════════════════════════════

_AXIS_STYLE = dict(gridcolor="#e2e8f0", linecolor="#cbd5e1", zerolinecolor="#e2e8f0")

_BASE = dict(
    plot_bgcolor="#fafbfc",
    paper_bgcolor="#ffffff",
    font=dict(family="Plus Jakarta Sans, Inter, sans-serif", size=12, color="#0f172a"),
    margin=dict(l=10, r=10, t=50, b=10),
    height=300,
)


def _chart_layout(*, xaxis_title=None, yaxis_title=None, xaxis=None, yaxis=None, **kwargs):
    """Shared Plotly layout; merges axis styling without duplicate yaxis/xaxis keys."""
    xa = dict(_AXIS_STYLE)
    ya = dict(_AXIS_STYLE)
    if xaxis:
        xa.update(xaxis)
    if yaxis:
        ya.update(yaxis)
    if xaxis_title:
        xa["title"] = xaxis_title
    if yaxis_title:
        ya["title"] = yaxis_title
    return {**_BASE, "xaxis": xa, "yaxis": ya, **kwargs}


def _cc(v):
    if v < 1.0:
        return "#ef4444"
    if v < 1.5:
        return "#f97316"
    if v < 2.0:
        return "#eab308"
    if v < 2.5:
        return "#10b981"
    return "#6366f1"


def chart_co_final(co_names, co_finals):
    avg_co = float(np.mean(co_finals)) if len(co_finals) else 0.0
    fig = go.Figure(go.Bar(
        x=co_names, y=np.round(co_finals, 3),
        marker_color=[_cc(v) for v in co_finals],
        text=np.round(co_finals, 3), textposition="outside",
    ))
    fig.add_hline(y=2.0, line_dash="dash", line_color="#ef4444",
                  annotation_text="Target 2.0", annotation_position="top right")
    fig.add_hline(y=avg_co, line_dash="dot", line_color="#6366f1",
                  annotation_text=f"CO Avg {avg_co:.3f}",
                  annotation_position="bottom right")
    fig.update_layout(**_chart_layout(
        title=dict(text="<b>Final CO Attainment</b>", font=dict(size=15, color="#0f172a"), x=0),
        xaxis_title="Course Outcomes",
        yaxis=dict(title="Attainment", range=[0, 3.6]),
    ))
    return fig


def chart_int_ext(co_names, co_int, co_ext, iw, uw, external_label="University"):
    int_v = np.where(np.isnan(co_int), 0, np.round(co_int, 3))
    ext_v = np.where(np.isnan(co_ext), 0, np.round(co_ext, 3))
    fig = go.Figure([
        go.Bar(name=f"Internal ({iw:.0f}%)", x=co_names, y=int_v,
               marker_color="#6366f1", text=int_v, textposition="outside"),
        go.Bar(name=f"{external_label} ({uw:.0f}%)", x=co_names, y=ext_v,
               marker_color="#06b6d4", text=ext_v, textposition="outside"),
    ])
    fig.update_layout(**_chart_layout(
        title=dict(text=f"<b>Internal vs {external_label} Attainment</b>",
                   font=dict(size=15, color="#0f172a"), x=0),
        barmode="group",
        xaxis_title="Course Outcomes",
        yaxis=dict(title="Attainment", range=[0, 3.6]),
        legend=dict(orientation="h", y=1.05, x=0),
    ))
    return fig


def chart_po(po_finals):
    po_names = [f"PO{j+1}" for j in range(NUM_POS)]
    fig = go.Figure(go.Bar(
        x=po_names, y=np.round(po_finals, 3),
        marker_color=[_cc(v) for v in po_finals],
        text=np.round(po_finals, 3), textposition="outside",
    ))
    fig.add_hline(y=2.0, line_dash="dash", line_color="#ef4444",
                  annotation_text="Target 2.0", annotation_position="top right")
    ymax = float(max(po_finals.max() * 1.35, 0.5)) if len(po_finals) else 0.5
    fig.update_layout(**_chart_layout(
        title=dict(text="<b>PO Attainment</b>", font=dict(size=15, color="#0f172a"), x=0),
        xaxis_title="Program Outcomes",
        yaxis=dict(title="Attainment", range=[0, ymax]),
    ))
    return fig


def chart_pso(pso_finals):
    pso_names = [f"PSO{j+1}" for j in range(NUM_PSOS)]
    fig = go.Figure(go.Bar(
        x=pso_names, y=np.round(pso_finals, 3),
        marker_color=["#6366f1", "#0d9488", "#d97706"][:NUM_PSOS],
        text=np.round(pso_finals, 3), textposition="outside", width=0.45,
    ))
    fig.add_hline(y=2.0, line_dash="dash", line_color="#ef4444",
                  annotation_text="Target 2.0", annotation_position="top right")
    ymax = float(max(pso_finals.max() * 1.35, 0.5)) if len(pso_finals) else 0.5
    fig.update_layout(**_chart_layout(
        title=dict(text="<b>PSO Attainment</b>", font=dict(size=15, color="#0f172a"), x=0),
        xaxis_title="Program Specific Outcomes",
        yaxis=dict(title="Attainment", range=[0, ymax]),
    ))
    return fig


def chart_di(exam_di_dict, threshold, all_exams):
    di_vals = [exam_di_dict[e]["DI"] for e in all_exams]
    fig = go.Figure(go.Bar(
        x=all_exams, y=di_vals,
        marker_color=["#10b981" if d >= threshold else "#ef4444" for d in di_vals],
        text=[f"{d:.4f}" for d in di_vals], textposition="outside",
    ))
    fig.add_hline(y=threshold, line_dash="dash", line_color="#f59e0b",
                  annotation_text=f"Threshold ({threshold:.2f})",
                  annotation_position="top right")
    ymax = max(max(di_vals, default=0.0) * 1.4, threshold * 2, 0.25)
    fig.update_layout(**_chart_layout(
        title=dict(text="<b>Discrimination Index per Exam</b>",
                   font=dict(size=15, color="#0f172a"), x=0),
        xaxis_title="Exam",
        yaxis=dict(title="DI", range=[0, ymax]),
    ))
    return fig


def chart_co_heatmap(co_names, co_finals, copo_matrix):
    po_names = [f"PO{j+1}" for j in range(NUM_POS)]
    contrib = np.array([[co_finals[i]/3*copo_matrix[i,j]
                         for j in range(NUM_POS)]
                        for i in range(NUM_COS)])
    fig = go.Figure(go.Heatmap(
        z=np.round(contrib, 3), x=po_names, y=co_names,
        colorscale=[[0,"#f8fafc"],[0.3,"#c7d2fe"],[0.7,"#6366f1"],[1,"#312e81"]],
        text=np.round(contrib, 3), texttemplate="%{text}",
        showscale=True, colorbar=dict(title="Contrib"),
    ))
    fig.update_layout(**_chart_layout(
        height=300,
        title=dict(text="<b>CO→PO Contribution Heatmap</b>",
                   font=dict(size=15, color="#0f172a"), x=0),
        xaxis_title="Program Outcomes",
        yaxis_title="Course Outcomes",
    ))
    return fig


# ══════════════════════════════════════════════════════════════════════════
#  SAMPLE EXCEL TEMPLATE
# ══════════════════════════════════════════════════════════════════════════

def generate_sample_excel(mode="Theory", components=None) -> BytesIO:
    np.random.seed(99)
    n = 60
    if mode == "Practical":
        components = components or []
        data = {
            "Student Name": [f"Student_{i+1:02d}" for i in range(n)],
            PRACTICAL_INTERNAL_EXAM: np.random.randint(8, 26, n),
        }
        for i, comp in enumerate(components):
            col = practical_component_col(i)
            upper = int(comp.get("max_marks", 20)) + 1
            data[col] = np.random.randint(max(0, upper // 3), max(1, upper), n)
        df = pd.DataFrame(data)
    else:
        df = pd.DataFrame({
            "Student Name": [f"Student_{i+1:02d}" for i in range(n)],
            "CCE 1":        np.random.randint(0, 11, n),
            "CCE 2":        np.random.randint(0, 11, n),
            "CCE 3":        np.random.randint(0, 11, n),
            "CCE 4":        np.random.randint(0, 11, n),
            "CCE 5":        np.random.randint(0, 11, n),
            "End Semester": np.random.randint(0, 51, n),
        })
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Student Marks")
        from openpyxl.styles import Font, PatternFill, Alignment
        ws = w.sheets["Student Marks"]
        for cell in ws[1]:
            cell.font      = Font(bold=True, color="FFFFFF", size=11)
            cell.fill      = PatternFill("solid", fgColor="1e1b4b")
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = \
                max(len(str(c.value or "")) for c in col) + 4
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════
#  SESSION STATE
# ══════════════════════════════════════════════════════════════════════════

def init_state():
    defs = {
        "processed":         False,
        "results":           {},
        "df":                None,
        "max_marks":         dict(DEFAULT_MAX),
        "subj_code":         DEFAULT_SUBJ_CODE,
        "subj_name":         DEFAULT_SUBJ_NAME,
        "co_stmts":          list(DEFAULT_CO_STMTS),
        "iw":                DEFAULT_INTERNAL_WEIGHT,
        "uw":                DEFAULT_UNIVERSITY_WEIGHT,
        "di_top_pct":        DEFAULT_DI_TOP_PCT,
        "di_bot_pct":        DEFAULT_DI_BOT_PCT,
        "di_threshold":      DEFAULT_DI_THRESHOLD,
        "mode":              DEFAULT_MODE,
        "practical_internal_max": DEFAULT_PRACTICAL_INTERNAL_MAX,
        "practical_components": [dict(c) for c in DEFAULT_PRACTICAL_COMPONENTS],
    }
    for k, v in defs.items():
        if k not in st.session_state:
            st.session_state[k] = v

    # Drop legacy BE theory column keys from max marks; ensure BTech exams exist
    _legacy_be_exams = {
        "Unit Test 1", "Unit Test 2", "Unit Test 3", "Prelim", "Insem", "Endsem",
    }
    mm = st.session_state.max_marks
    for _k in list(mm.keys()):
        if _k in _legacy_be_exams:
            del mm[_k]
    for _ex, _mv in DEFAULT_MAX.items():
        if _ex not in mm:
            mm[_ex] = _mv

    # Keep persisted state aligned with the configured CO / PO dimensions.
    co_stmts = list(st.session_state.get("co_stmts", []))
    if len(co_stmts) != NUM_COS:
        st.session_state.co_stmts = list(DEFAULT_CO_STMTS) if not co_stmts else (co_stmts + [""] * NUM_COS)[:NUM_COS]

    results = st.session_state.get("results")
    if isinstance(results, dict) and results:
        result_cos = len(results.get("co_names", []))
        result_pos = len(results.get("po_finals", []))
        if result_cos != NUM_COS or result_pos != NUM_POS:
            st.session_state.results = {}
            st.session_state.processed = False

init_state()


# ══════════════════════════════════════════════════════════════════════════
#  CALCULATION ORCHESTRATOR
# ══════════════════════════════════════════════════════════════════════════

def run_calculations(df, max_marks, copo_mat, copso_mat, coexam_mat,
                     subj_code, co_stmts, iw, uw,
                     di_top_pct, di_bot_pct, mode="Theory"):
    try:
        co_names = [f"CO{i+1}" for i in range(NUM_COS)]
        ir, ur   = normalize_weights(iw, uw)
        _, ext_label = assessment_labels(mode)

        internal_exams, external_exams, all_exams = get_exam_sets(mode)

        exam_att_dict = {}
        exam_di_dict  = {}
        for exam in all_exams:
            marks = df[exam].fillna(0).values
            att, _ = exam_attainment(marks, max_marks[exam])
            DI, H, L, k_top, k_bot = discrimination_index(marks, max_marks[exam], di_top_pct, di_bot_pct)
            exam_att_dict[exam] = att
            exam_di_dict[exam]  = {"DI": DI, "H": H, "L": L,
                                   "k_top": k_top, "k_bot": k_bot, "N": len(marks)}

        co_int_att, co_ext_att = compute_co_attainments(
            exam_att_dict, coexam_mat, internal_exams, external_exams, all_exams
        )
        co_finals = compute_co_finals(co_int_att, co_ext_att, ir, ur)
        po_contrib,  po_finals  = compute_po_pso(co_finals, copo_mat)
        pso_contrib, pso_finals = compute_po_pso(co_finals, copso_mat)

        st.session_state.results = {
            "co_names":      co_names,
            "co_stmts":      co_stmts,
            "subj_code":     subj_code,
            "exam_att_dict": exam_att_dict,
            "exam_di_dict":  exam_di_dict,
            "co_int_att":    co_int_att,
            "co_ext_att":    co_ext_att,
            "co_finals":     co_finals,
            "po_contrib":    po_contrib,
            "po_finals":     po_finals,
            "pso_contrib":   pso_contrib,
            "pso_finals":    pso_finals,
            "copo_matrix":   copo_mat,
            "copso_matrix":  copso_mat,
            "coexam_matrix": coexam_mat,
            "max_marks":     max_marks,
            "iw": iw, "uw": uw,
            "df": df,
            "mode": mode,
            "external_label": ext_label,
            "internal_exams": internal_exams,
            "external_exams": external_exams,
            "all_exams": all_exams,
        }
        st.session_state.processed = True

    except Exception as e:
        import traceback
        st.error(f"Calculation could not be completed: {e}")
        with st.expander("Technical details", expanded=False):
            st.code(traceback.format_exc())
        st.session_state.processed = False


# ══════════════════════════════════════════════════════════════════════════
#  MAIN APP
# ══════════════════════════════════════════════════════════════════════════

def run_app():
    iw = st.session_state.iw
    uw = st.session_state.uw
    mode = st.session_state.mode
    internal_exams, external_exams, _ = get_exam_sets(mode)
    _, ext_label = assessment_labels(mode)

    # ── HEADER ─────────────────────────────────────────────────────────
    st.markdown(f"""
    <div class="nba-hdr">
      <div class="hdr-logo">OBE</div>
      <div class="hdr-text">
        <div class="hdr-title">BTech CO-PO Attainment &amp; Academic Analytics System</div>
        <div class="hdr-sub">BTech Theory Pattern · CCE Based Internal Assessment · NBA Accreditation Tool</div>
      </div>
      <div class="hdr-pills">
        <div class="hdr-pill">BTech</div>
        <div class="hdr-pill">{len(internal_exams)} Internal</div>
        <div class="hdr-pill">{len(external_exams)} {ext_label}</div>
        <div class="hdr-pill">{iw:.0f}:{uw:.0f} Weightage</div>
        <div class="hdr-pill">Mode: {mode}</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # ── TABS ───────────────────────────────────────────────────────────
    t_in, t_res, t_graph, t_exp = st.tabs([
        "Input & Configuration",
        "Results & Reports",
        "Analytics Charts",
        "Export",
    ])

    # ╔════════════════════════════════════════╗
    # ║  TAB 1 — INPUT                        ║
    # ╚════════════════════════════════════════╝
    with t_in:
        with st.container(border=True):
            ui_stitle("Assessment Mode")
            mode = st.radio(
                "Mode",
                ["Theory", "Practical"],
                horizontal=True,
                index=0 if st.session_state.mode == "Theory" else 1,
                label_visibility="collapsed",
            )
            st.session_state.mode = mode
        _, ext_label = assessment_labels(mode)

        practical_components = st.session_state.practical_components
        if mode == "Practical":
            with st.container(border=True):
                ui_stitle("Practical Configuration")
                cpi, cpa = st.columns([1, 1], vertical_alignment="bottom")
                with cpi:
                    st.session_state.practical_internal_max = st.number_input(
                        "Internal Max Marks",
                        min_value=1,
                        max_value=500,
                        value=int(st.session_state.practical_internal_max),
                        key="prac_internal_max_input",
                    )
                with cpa:
                    if st.button("Add External Component", type="secondary", key="add_practical_component"):
                        practical_components.append({"name": f"Component {len(practical_components)+1}", "max_marks": 25})

                for idx, comp in enumerate(practical_components):
                    cx1, cx2, cx3 = st.columns([2, 1, 0.6], vertical_alignment="bottom")
                    with cx1:
                        comp["name"] = st.text_input("Component Name", value=comp.get("name", ""), key=f"pc_name_{idx}")
                    with cx2:
                        comp["max_marks"] = st.number_input("Max Marks", min_value=1, max_value=500,
                                                            value=int(comp.get("max_marks", 25)), key=f"pc_mm_{idx}")
                    with cx3:
                        if st.button("Remove", type="secondary", key=f"pc_del_{idx}") and len(practical_components) > 1:
                            practical_components.pop(idx)
                            st.rerun()
                st.session_state.practical_components = practical_components
                st.markdown('<div class="ainfo">External attainment in Practical mode is computed by combining all configured external components per student.</div>', unsafe_allow_html=True)

        internal_exams, external_exams, all_exams = get_exam_sets(mode)
        col_L, col_R = st.columns([1.08, 0.92], gap="small")
        setup_pills = [
            ("neutral", f"{mode} mode active"),
            ("neutral", f"{len(all_exams)} exam column(s) expected"),
            ("neutral", f"{len(practical_components) if mode == 'Practical' else len(external_exams)} external component(s)"),
        ]
        setup_pills_html = "".join(f'<span class="ready-pill {tone}">{label}</span>' for tone, label in setup_pills)
        st.markdown(f'<div class="ready-strip">{setup_pills_html}</div>', unsafe_allow_html=True)

        # ── LEFT COLUMN ───────────────────────────────────────────────
        with col_L:
            ui_section("Course and Dataset", "Define the subject, CO statements, and upload the marks workbook.")
            ui_stitle("Subject Information")
            sc1, sc2 = st.columns([1, 2])
            with sc1:
                subj_code = st.text_input("Subject Code", value=st.session_state.subj_code, key="inp_code")
            with sc2:
                subj_name = st.text_input("Subject Name", value=st.session_state.subj_name, key="inp_name")
            st.session_state.subj_code = subj_code
            st.session_state.subj_name = subj_name

            ui_stitle("Course Outcomes")
            co_names_cur = [f"CO{i+1}" for i in range(NUM_COS)]
            co_stmts = []
            for i in range(NUM_COS):
                stmt = st.text_input(
                    label=co_names_cur[i],
                    value=st.session_state.co_stmts[i] if i < len(st.session_state.co_stmts) else "",
                    key=f"co_stmt_{i}",
                )
                co_stmts.append(stmt)
            st.session_state.co_stmts = co_stmts

            ui_stitle("Upload Student Marks")
            req_cols = ["Student Name"] + (all_exams if mode == "Theory" else [PRACTICAL_INTERNAL_EXAM] + practical_component_columns(practical_components))
            st.markdown(f'<div class="ainfo"><b>Required columns:</b> {", ".join(req_cols)}</div>',
                        unsafe_allow_html=True)
            if mode == "Theory":
                st.markdown(
                    '<div class="ainfo" style="margin-top:6px"><b>Theory marks validation:</b> '
                    "Each assessment is clipped to the configured maximum (defaults: CCE 1–5 → 10; End Semester → 50).</div>",
                    unsafe_allow_html=True,
                )
            if mode == "Practical":
                comp_note = ", ".join([f"{practical_component_col(i)} ({c['name']})" for i, c in enumerate(practical_components)])
                st.markdown(f'<div class="ainfo"><b>External components:</b> {comp_note}</div>', unsafe_allow_html=True)

            sample_bytes = generate_sample_excel(mode=mode, components=practical_components).getvalue()
            st.download_button("Download Sample Template", data=sample_bytes,
                               file_name=f"sample_marks_{subj_code}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               key="download_sample_template",
                               on_click="ignore",
                               width="stretch")

            uploaded = st.file_uploader("Upload .xlsx / .xls file", type=["xlsx", "xls"])
            if uploaded:
                try:
                    df_raw = pd.read_excel(uploaded)
                    expected = all_exams if mode == "Theory" else [PRACTICAL_INTERNAL_EXAM] + practical_component_columns(practical_components)
                    missing = [c for c in expected if c not in df_raw.columns]
                    if missing:
                        st.error(f"Missing columns: {', '.join(missing)}")
                    else:
                        st.session_state.df = df_raw
                        st.success(f"{len(df_raw)} students loaded successfully.")
                        with st.expander("Preview (first 10 rows)"):
                            st.dataframe(df_raw.head(10), width="stretch")
                except Exception as e:
                    st.error(f"Error reading file: {e}")

            ui_stitle("Maximum Marks per Exam")
            c1, c2, c3 = st.columns(3)
            mm = {}
            mm_source = all_exams if mode == "Theory" else [PRACTICAL_INTERNAL_EXAM] + practical_component_columns(practical_components)
            for idx, exam in enumerate(mm_source):
                with [c1, c2, c3][idx % 3]:
                    default_val = DEFAULT_MAX.get(exam, 100)
                    if mode == "Practical" and exam == PRACTICAL_INTERNAL_EXAM:
                        default_val = st.session_state.practical_internal_max
                    if mode == "Practical" and exam.startswith("External_Component_"):
                        comp_idx = int(exam.split("_")[-1]) - 1
                        if comp_idx < len(practical_components):
                            default_val = practical_components[comp_idx]["max_marks"]
                    mm[exam] = st.number_input(exam, min_value=1, max_value=500,
                                               value=int(st.session_state.max_marks.get(exam, default_val)),
                                               key=f"mm_{exam}")
            if mode == "Practical":
                mm[PRACTICAL_EXTERNAL_EXAM] = sum(mm[c] for c in practical_component_columns(practical_components))
            st.session_state.max_marks = mm

            # Attainment table
            ui_stitle("Attainment Level Reference")
            st.markdown("""
            <div class="subtable-wrap">
              <div class="nba-wrap">
                <table class="nba-tbl">
                  <tr><th>Level</th><th>Percentage</th><th>Meaning</th></tr>
                  <tr><td>0</td><td>&lt; 40%</td><td style="color:#b91c1c;font-weight:700">Not Achieved</td></tr>
                  <tr><td>1</td><td>40%–55%</td><td style="color:#92400e;font-weight:700">Partially Achieved</td></tr>
                  <tr><td>2</td><td>56%–70%</td><td style="color:#065f46;font-weight:700">Achieved</td></tr>
                  <tr><td>3</td><td>&gt; 70%</td><td style="color:#1e3a8a;font-weight:700">Highly Achieved</td></tr>
                </table>
              </div>
            </div>""", unsafe_allow_html=True)

        # ── RIGHT COLUMN ──────────────────────────────────────────────
        with col_R:
            ui_section("Evaluation Rules and Mapping", "Configure weightages, discrimination index, and mappings.")
            ui_stitle("Assessment Weightages")
            w1, w2 = st.columns(2)
            with w1:
                iw = st.number_input("Internal Weight (%)", 0.0, 100.0,
                                     float(st.session_state.iw), 5.0, key="iw_inp")
            with w2:
                uw = st.number_input(f"{ext_label} Weight (%)", 0.0, 100.0,
                                     float(st.session_state.uw), 5.0, key="uw_inp")
            st.session_state.iw = iw
            st.session_state.uw = uw
            ir, ur = normalize_weights(iw, uw)
            if abs(iw + uw - 100) > 0.01:
                st.markdown(f'<div class="ainfo">Total = {iw+uw:.1f}%. Will be normalized to {ir*100:.1f}% / {ur*100:.1f}%.</div>',
                            unsafe_allow_html=True)

            # ── DI CONFIGURATION ──────────────────────────────────────
            ui_stitle("Discrimination Index Settings")
            st.markdown("""
            <div class="di-config">
              <b>Configure discrimination index parameters</b><br>
              <span style="font-size:12px;color:#555">
                DI = (Mean of top group − Mean of bottom group) / Max Marks<br>
                Adjust group sizes and threshold as per your institutional norms.
              </span>
            </div>""", unsafe_allow_html=True)
            d1, d2, d3 = st.columns(3)
            with d1:
                di_top = st.number_input("Top Group (%)", 5.0, 50.0,
                                         float(st.session_state.di_top_pct), 1.0,
                                         key="di_top",
                                         help="Top N% students used for H (high group)")
            with d2:
                di_bot = st.number_input("Bottom Group (%)", 5.0, 50.0,
                                         float(st.session_state.di_bot_pct), 1.0,
                                         key="di_bot",
                                         help="Bottom N% students used for L (low group)")
            with d3:
                di_thr = st.number_input("DI Threshold", 0.01, 1.0,
                                         float(st.session_state.di_threshold), 0.01,
                                         format="%.2f",
                                         key="di_thr",
                                         help="DI ≥ threshold = Good; DI < threshold = Needs revision")
            st.session_state.di_top_pct   = di_top
            st.session_state.di_bot_pct   = di_bot
            st.session_state.di_threshold = di_thr

            # ── MAPPINGS ──────────────────────────────────────────────
            po_cols  = [f"PO{j+1}"  for j in range(NUM_POS)]
            pso_cols = [f"PSO{j+1}" for j in range(NUM_PSOS)]

            ui_stitle("CO–Exam Mapping (1 = assessed, 0 = not assessed)")
            if mode == "Theory":
                df_coexam_init = pd.DataFrame(DEFAULT_COEXAM, index=co_names_cur, columns=all_exams)
            else:
                df_coexam_init = pd.DataFrame(np.ones((NUM_COS, len(all_exams))), index=co_names_cur, columns=all_exams)
            edited_coexam  = st.data_editor(df_coexam_init, key=f"coexam_ed_{mode}_{NUM_COS}",
                                            width="stretch", num_rows="fixed",
                                            column_config={e: st.column_config.NumberColumn(e, min_value=0, max_value=1, step=1) for e in all_exams})

            ui_stitle("CO–PO Mapping (0 = None to 3 = High)")
            df_copo_init  = pd.DataFrame(DEFAULT_COPO, index=co_names_cur, columns=po_cols)
            edited_copo   = st.data_editor(df_copo_init, key=f"copo_ed_{NUM_COS}_{NUM_POS}",
                                           width="stretch", num_rows="fixed",
                                           column_config={c: st.column_config.NumberColumn(c, min_value=0, max_value=3, step=1) for c in po_cols})

            ui_stitle("CO–PSO Mapping (0 = None to 3 = High)")
            df_copso_init = pd.DataFrame(DEFAULT_COPSO, index=co_names_cur, columns=pso_cols)
            edited_copso  = st.data_editor(df_copso_init, key=f"copso_ed_{NUM_COS}_{NUM_PSOS}",
                                           width="stretch", num_rows="fixed",
                                           column_config={c: st.column_config.NumberColumn(c, min_value=0, max_value=3, step=1) for c in pso_cols})

            st.markdown(f"""
            <div class="ainfo" style="margin-top:14px">
              <b>Formula reference</b><br>
              Exam CO Att = (0·P + 1·Q + 2·R + 3·S) / N<br>
              CO Internal = avg of mapped internal exams per CO<br>
              CO {ext_label} = avg of mapped {ext_label.lower()} exams per CO<br>
              Final CO = {ir*100:.2f}% × Internal + {ur*100:.2f}% × {ext_label} (normalized from {iw:.1f}% / {uw:.1f}%); if only internal or only external is mapped for a CO, Final CO uses the available part only<br>
              PO<sub>j</sub> = avg<sub>all COs</sub>[ (CO<sub>i</sub>/3) × w<sub>ij</sub> ]<br>
              DI = (H − L) / Max Marks &nbsp;[top {di_top:.0f}% vs bottom {di_bot:.0f}%]
            </div>
            """, unsafe_allow_html=True)

        # ── CALCULATE ─────────────────────────────────────────────────
        with st.container(border=True):
            ui_section("Run the Analysis", "Review readiness badges, then calculate attainment.")
            calc = st.button("Calculate CO-PO Attainment", type="primary", width="stretch")

            students_loaded = 0 if st.session_state.df is None else len(st.session_state.df)
            co_filled = sum(1 for stmt in co_stmts if str(stmt).strip())
            ready_pills = [
                ("ok" if students_loaded > 0 else "warn", f"{students_loaded} student record(s) loaded" if students_loaded > 0 else "Marks file pending"),
                ("ok" if co_filled == NUM_COS else "warn", f"{co_filled}/{NUM_COS} CO statements filled"),
                ("ok" if abs(iw + uw - 100) <= 0.01 else "neutral", f"Weight total {iw + uw:.1f}%"),
                ("neutral", f"{len(all_exams)} exam column(s) configured"),
            ]
            pills_html = "".join(f'<span class="ready-pill {tone}">{label}</span>' for tone, label in ready_pills)
            st.markdown(f'<div class="ready-strip">{pills_html}</div>', unsafe_allow_html=True)

        if calc:
            if st.session_state.df is None:
                st.error("Please upload a student marks Excel file first.")
            else:
                # clean mappings
                def clean_df(df, lo, hi):
                    df2 = df.copy()
                    for col in df2.columns:
                        df2[col] = pd.to_numeric(df2[col], errors="coerce").fillna(0).clip(lo, hi)
                    return df2
                coexam_clean = clean_df(edited_coexam, 0, 1)
                copo_clean   = clean_df(edited_copo,   0, 3)
                copso_clean  = clean_df(edited_copso,  0, 3)

                with st.spinner("Computing attainment…"):
                    df_calc = st.session_state.df.copy()
                    max_marks_calc = dict(st.session_state.max_marks)
                    if mode == "Practical":
                        component_cols = practical_component_columns(practical_components)
                        df_calc = build_practical_calc_df(df_calc, component_cols)
                        max_marks_calc = {
                            PRACTICAL_INTERNAL_EXAM: float(st.session_state.max_marks[PRACTICAL_INTERNAL_EXAM]),
                            PRACTICAL_EXTERNAL_EXAM: float(sum(st.session_state.max_marks[c] for c in component_cols)),
                        }
                    else:
                        _, _, theory_exams = get_exam_sets("Theory")
                        for ex in theory_exams:
                            if ex in df_calc.columns:
                                cap = float(max_marks_calc.get(ex, DEFAULT_MAX.get(ex, 1)))
                                df_calc[ex] = pd.to_numeric(df_calc[ex], errors="coerce").fillna(0).clip(lower=0, upper=cap)
                    run_calculations(
                        df_calc,
                        max_marks_calc,
                        copo_clean.values.astype(float),
                        copso_clean.values.astype(float),
                        coexam_clean.values.astype(float),
                        subj_code, co_stmts,
                        iw, uw, di_top, di_bot, mode,
                    )
                if st.session_state.processed:
                    st.success("Calculation complete. Open Results, Analytics Charts, or Export.")

    # ╔════════════════════════════════════════╗
    # ║  TAB 2 — RESULTS                      ║
    # ╚════════════════════════════════════════╝
    with t_res:
        if not st.session_state.processed:
            st.markdown('<div class="empty-state"><div class="empty-title">Results will appear here after calculation</div><div class="empty-copy">Complete the input setup, upload the marks workbook, and run the attainment engine once. This tab is intended to become the presentation-ready review space for your course outcome analysis.</div></div>',
                        unsafe_allow_html=True)
        else:
            R    = st.session_state.results
            co_n = R["co_names"]
            di_t = st.session_state.di_threshold
            top_p = st.session_state.di_top_pct
            bot_p = st.session_state.di_bot_pct

            # Subject box
            co_items = "".join(
                f'<div class="co-item"><span class="co-lbl">{n}:</span> {s}</div>'
                for n, s in zip(co_n, R["co_stmts"])
            )
            st.markdown(f"""
            <div class="subj-box">
              <div class="sc">Subject Profile</div>
              <div class="sn">{R["subj_code"]} &nbsp;–&nbsp; {st.session_state.subj_name}</div>
              <div class="co-grid">{co_items}</div>
            </div>""", unsafe_allow_html=True)

            # Metric strip
            co_ok  = int((R["co_finals"] >= 2.0).sum())
            avg_co = float(R["co_finals"].mean())
            po_v   = R["po_finals"]
            avg_po = float(po_v[po_v > 0].mean()) if (po_v > 0).any() else 0.0
            low_di = sum(1 for d in R["exam_di_dict"].values() if d["DI"] < di_t)
            n_stu  = len(R["df"])

            st.markdown(f"""
            <div class="mstrip">
              <div class="mbox" style="--ac:#6366f1">
                <div class="ml">Avg CO Attainment</div>
                <div class="mv">{avg_co:.3f}</div>
                <div class="ms">Out of 3.0</div>
              </div>
              <div class="mbox" style="--ac:#10b981">
                <div class="ml">COs at Target ≥2.0</div>
                <div class="mv">{co_ok} / {NUM_COS}</div>
                <div class="ms">Threshold met</div>
              </div>
              <div class="mbox" style="--ac:#06b6d4">
                <div class="ml">Avg PO Attainment</div>
                <div class="mv">{avg_po:.3f}</div>
                <div class="ms">Mapped POs only</div>
              </div>
              <div class="mbox" style="--ac:{'#10b981' if low_di == 0 else '#ef4444'}">
                <div class="ml">Low DI Exams</div>
                <div class="mv">{low_di} / {len(R["all_exams"])}</div>
                <div class="ms">DI &lt; {di_t:.2f}</div>
              </div>
              <div class="mbox" style="--ac:#8b5cf6">
                <div class="ml">Total Students</div>
                <div class="mv">{n_stu}</div>
                <div class="ms">Analysed</div>
              </div>
            </div>
            """, unsafe_allow_html=True)

            ui_stitle("CO Attainment Overview")
            st.markdown('<div class="tbl-label">Internal assessment</div>', unsafe_allow_html=True)
            st.markdown(render_internal_table(co_n, R["coexam_matrix"],
                                              R["exam_att_dict"], R["co_int_att"],
                                              R["internal_exams"], R["all_exams"]),
                        unsafe_allow_html=True)

            st.markdown(
                f'<div class="tbl-label">{R.get("external_label", "University")} assessment</div>',
                unsafe_allow_html=True,
            )
            st.markdown(render_external_table(co_n, R["coexam_matrix"],
                                              R["exam_att_dict"], R["co_ext_att"],
                                              R["external_exams"], R["all_exams"]),
                        unsafe_allow_html=True)

            st.markdown(f"""
            <div class="ainfo" style="margin:4px 0">
              <b>Actual CO Attainment = {R['uw']:.1f}% {R.get('external_label', 'University')} + {R['iw']:.1f}% Internal</b>
            </div>""", unsafe_allow_html=True)
            st.markdown('<div class="tbl-label">Weighted final CO attainment</div>', unsafe_allow_html=True)
            st.markdown(render_final_table(co_n, R["co_int_att"], R["co_ext_att"],
                                           R["co_finals"], R["iw"], R["uw"], R.get("external_label", "University")),
                        unsafe_allow_html=True)
            # ── DI ───────────────────────────────────────────────────
            st.markdown(
                f'<div class="stitle">Discrimination Index '
                f'<span class="stitle-meta">Top {top_p:.0f}% / Bottom {bot_p:.0f}% · Threshold {di_t:.2f}</span></div>',
                unsafe_allow_html=True,
            )
            st.markdown(render_di_table(R["exam_di_dict"], top_p, bot_p, di_t, R["internal_exams"]),
                        unsafe_allow_html=True)
            ui_stitle("PO / PSO Attainment Overview")
            st.markdown(
                '<div class="tbl-label">PO attainment '
                '<span style="font-weight:500;color:var(--muted)">'
                '(contrib = CO att / 3 × weight; averaged over COs)</span></div>',
                unsafe_allow_html=True,
            )
            st.markdown(render_po_table(co_n, R["co_finals"],
                                        R["po_contrib"], R["po_finals"]),
                        unsafe_allow_html=True)

            st.markdown('<div class="tbl-label">PSO attainment</div>', unsafe_allow_html=True)
            st.markdown(render_pso_table(co_n, R["co_finals"],
                                         R["pso_contrib"], R["pso_finals"]),
                        unsafe_allow_html=True)
            ui_stitle("Insights & Recommendations")
            any_issue = False

            for co, val in zip(co_n, R["co_finals"]):
                if val < 2.0:
                    any_issue = True
                    st.markdown(f"""
                    <div class="awarn">
                      <b>{co} — Attainment LOW ({val:.4f})</b><br>
                      • Revise teaching strategy for this CO<br>
                      • Conduct remedial sessions / extra tutorials<br>
                      • Align exam questions to Bloom's taxonomy
                    </div>""", unsafe_allow_html=True)

            for exam, d in R["exam_di_dict"].items():
                if d["DI"] < di_t:
                    any_issue = True
                    st.markdown(f"""
                    <div class="awarn">
                      <b>{exam} — Low DI = {d['DI']:.4f} (threshold {di_t:.2f})</b><br>
                      • Question paper does not differentiate well<br>
                      • Rebalance difficulty: Easy 30% / Medium 50% / Hard 20%<br>
                      • Add Higher Order Thinking (HOT) questions
                    </div>""", unsafe_allow_html=True)

            good_cos = [(n, v) for n, v in zip(co_n, R["co_finals"]) if v >= 2.5]
            if good_cos:
                st.markdown(f"""
                <div class="agood">
                  <b>High-performing COs:</b> {', '.join(f'{n} ({v:.3f})' for n, v in good_cos)}<br>
                  Excellent outcomes for these COs — maintain current approach.
                </div>""", unsafe_allow_html=True)

            if not any_issue:
                st.markdown("""
                <div class="agood">
                  <b>All CO targets (≥ 2.0) met and DI values healthy!</b><br>
                  Course is on track for NBA accreditation.
                </div>""", unsafe_allow_html=True)

            st.markdown(f"""
            <div class="ainfo">
              <b>Formula notes</b> &nbsp;
              Exam CO Att = (0·P + 1·Q + 2·R + 3·S) / N &nbsp;|&nbsp;
              PO/PSO = avg<sub>all COs</sub>[(CO/3) × w] &nbsp;|&nbsp;
              DI = (H − L) / Max Marks with top {top_p:.0f}% / bottom {bot_p:.0f}% groups
            </div>""", unsafe_allow_html=True)
    # ╔════════════════════════════════════════╗
    # ║  TAB 3 — CHARTS                       ║
    # ╚════════════════════════════════════════╝
    with t_graph:
        if not st.session_state.processed:
            st.markdown('<div class="empty-state"><div class="empty-title">Charts are waiting for analysed data</div><div class="empty-copy">After the attainment run completes, this tab will show cleaner visual comparisons for CO, PO, PSO, discrimination index, and mapping impact.</div></div>', unsafe_allow_html=True)
        else:
            R    = st.session_state.results
            co_n = R["co_names"]
            di_t = st.session_state.di_threshold
            ui_section(
                "Performance Visuals",
                "Side-by-side charts for academic review and accreditation discussions.",
            )

            g1, g2 = st.columns(2, gap="medium")
            with g1:
                st.plotly_chart(chart_co_final(co_n, R["co_finals"]),
                                width="stretch")
            with g2:
                st.plotly_chart(chart_int_ext(co_n, R["co_int_att"],
                                              R["co_ext_att"], R["iw"], R["uw"], R.get("external_label", "University")),
                                width="stretch")

            g3, g4 = st.columns(2, gap="medium")
            with g3:
                st.plotly_chart(chart_po(R["po_finals"]), width="stretch")
            with g4:
                st.plotly_chart(chart_pso(R["pso_finals"]), width="stretch")

            g5, g6 = st.columns(2, gap="medium")
            with g5:
                st.plotly_chart(chart_di(R["exam_di_dict"], di_t, R["all_exams"]),
                                width="stretch")
            with g6:
                st.plotly_chart(chart_co_heatmap(co_n, R["co_finals"], R["copo_matrix"]),
                                width="stretch")

    # ╔════════════════════════════════════════╗
    # ║  TAB 4 — EXPORT                       ║
    # ╚════════════════════════════════════════╝
    with t_exp:
        if not st.session_state.processed:
            st.markdown('<div class="empty-state"><div class="empty-title">Export becomes available after a successful run</div><div class="empty-copy">Generate the attainment results first. Then this tab will provide polished Excel and PDF outputs together with a quick preview of the final reporting summary.</div></div>', unsafe_allow_html=True)
        else:
            R    = st.session_state.results
            top_p = st.session_state.di_top_pct
            bot_p = st.session_state.di_bot_pct
            di_t  = st.session_state.di_threshold

            st.markdown('<div class="download-grid"><div class="download-card"><div class="dk">Workbook export</div><div class="dv">Excel report</div><div class="ds">Multi-sheet output for data review, archival, and committee sharing.</div></div><div class="download-card"><div class="dk">Printable export</div><div class="dv">PDF report</div><div class="ds">A polished A4 summary for presentation, documentation, and accreditation evidence.</div></div></div>', unsafe_allow_html=True)
            ui_stitle("Download Reports")
            st.markdown("""
            <div class="ainfo">
              <b>Excel Report (7 Sheets):</b> Cover · Student Data · CO Attainment (a/b/c) · DI · PO · PSO · Summary<br>
              <b>PDF Report (A4 Portrait):</b> BTech assessment overview · CO statements · Student summary · Attainment tables · DI · PO/PSO · Summary · Charts (no radar)<br>
              <span style="font-size:12px;color:#4b5563">Theory mode uses CCE 1–5 plus End Semester; exports follow your configured labels and weights.</span>
            </div>""", unsafe_allow_html=True)

            with st.spinner("Building Excel + PDF reports…"):
                excel_buf = export_excel(R, st.session_state.subj_name, top_p, bot_p, di_t)
                pdf_buf   = export_pdf(R, st.session_state.subj_name, top_p, bot_p, di_t)
                excel_bytes = excel_buf.getvalue()
                pdf_bytes   = pdf_buf.getvalue()

            d1, d2 = st.columns(2, gap="medium")
            with d1:
                st.download_button(
                    "Download Excel Report (.xlsx)",
                    data=excel_bytes,
                    file_name=f"CO_PO_Report_{R['subj_code']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="download_excel_report",
                    on_click="ignore",
                    width="stretch",
                )
            with d2:
                st.download_button(
                    "Download PDF Report (A4)",
                    data=pdf_bytes,
                    file_name=f"CO_PO_Report_{R['subj_code']}.pdf",
                    mime="application/pdf",
                    key="download_pdf_report",
                    on_click="ignore",
                    width="stretch",
                )

            ui_stitle("Quick Summary Preview")
            st.markdown('<div class="preview-note">Use this preview as a final confidence check before downloading or circulating the generated reports.</div>', unsafe_allow_html=True)

            p1, p2, p3 = st.columns(3, gap="medium")
            with p1:
                st.markdown("**CO Attainment Summary**")
                co_df = pd.DataFrame({
                    "CO":          R["co_names"],
                    "Internal":    ["—" if np.isnan(v) else f"{v:.4f}" for v in R["co_int_att"]],
                    R.get("external_label", "University"):  ["—" if np.isnan(v) else f"{v:.4f}" for v in R["co_ext_att"]],
                    "Final":       np.round(R["co_finals"], 4),
                    "Target":      [_attainment_status_met(v) for v in R["co_finals"]],
                })
                st.dataframe(co_df, width="stretch", hide_index=True)
            with p2:
                st.markdown("**PO Attainment**")
                po_df = pd.DataFrame({
                    "PO":          [f"PO{j+1}" for j in range(NUM_POS)],
                    "Attainment":  np.round(R["po_finals"], 4),
                    "Status":      [_attainment_status_short(v) for v in R["po_finals"]],
                })
                st.dataframe(po_df, width="stretch", hide_index=True)
            with p3:
                st.markdown("**PSO Attainment**")
                pso_df = pd.DataFrame({
                    "PSO":         [f"PSO{j+1}" for j in range(NUM_PSOS)],
                    "Attainment":  np.round(R["pso_finals"], 4),
                    "Status":      [_attainment_status_short(v) for v in R["pso_finals"]],
                })
                st.dataframe(pso_df, width="stretch", hide_index=True)


# ══════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    run_app()
